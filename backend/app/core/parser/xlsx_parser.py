"""XLSX 表格解析器

使用 openpyxl 从 Excel 文件中提取工作表数据。
用于标书编写功能中的素材文档(报价表/人员表/业绩表等)解析。
"""

from dataclasses import dataclass, field
from typing import Optional

from loguru import logger


@dataclass
class XlsxSheet:
    """单个工作表数据"""
    name: str
    rows: list[list[str]]
    header_row: Optional[list[str]] = None
    row_count: int = 0
    col_count: int = 0


@dataclass
class ParsedXlsx:
    """XLSX 解析结果"""
    full_text: str
    sheets: list[XlsxSheet] = field(default_factory=list)


class XlsxParser:
    """XLSX 表格解析器"""

    def parse(self, file_path: str) -> ParsedXlsx:
        """解析 XLSX 文件

        Args:
            file_path: XLSX 文件路径

        Returns:
            ParsedXlsx 解析结果
        """
        try:
            from openpyxl import load_workbook
        except ImportError:
            logger.error("openpyxl 未安装，无法解析 XLSX 文件")
            return ParsedXlsx(full_text="")

        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
        except Exception as e:
            logger.error(f"XLSX 文件打开失败: {e}")
            return ParsedXlsx(full_text="")

        sheets: list[XlsxSheet] = []
        text_parts: list[str] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = self._extract_rows(ws)
            if not rows:
                continue

            header = rows[0] if rows else None
            sheet = XlsxSheet(
                name=sheet_name,
                rows=rows,
                header_row=header,
                row_count=len(rows),
                col_count=len(rows[0]) if rows else 0,
            )
            sheets.append(sheet)

            text_parts.append(f"[工作表: {sheet_name}]")
            text_parts.append(self._sheet_to_text(sheet))

        wb.close()

        full_text = "\n\n".join(text_parts)
        logger.info(
            f"XLSX 解析完成: sheets={len(sheets)}, "
            f"total_rows={sum(s.row_count for s in sheets)}"
        )
        return ParsedXlsx(full_text=full_text, sheets=sheets)

    def _extract_rows(self, ws: object) -> list[list[str]]:
        """提取工作表行数据"""
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(cell) if cell is not None else "" for cell in row]
            if any(c.strip() for c in cells):
                rows.append(cells)
        return rows

    def _sheet_to_text(self, sheet: XlsxSheet) -> str:
        """将工作表转为文本"""
        lines: list[str] = []
        for row in sheet.rows:
            lines.append(" | ".join(row))
        return "\n".join(lines)

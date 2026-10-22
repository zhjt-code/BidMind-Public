"""生成图示与正文/渲染管道的桥接。

职责: 维护 <figure data-diagram-id> 协议 -- 漏插兜底、占位重写、资产解析与插入。
不做: 不编译 IR、不渲染 PNG、不处理素材图片(data-image-id)。
"""

from __future__ import annotations

import re
from typing import TYPE_CHECKING, Iterable

from loguru import logger

from app.core.xml_compat import sanitize_xml_compatible_text

if TYPE_CHECKING:
    from app.core.writer.visuals.materializer import DiagramAsset

DIAGRAM_FIGURE_RE = re.compile(
    r'<figure\s+data-diagram-id="(\d+)"[^>]*>.*?</figure>',
    re.DOTALL,
)
DIAGRAM_PLACEHOLDER_TEXT = "[生成图示占位]"
_DIAGRAM_PLACEHOLDER_RE = re.compile(
    r'<p class="__bid_diagram__" data-diagram-id="(\d+)">\[生成图示占位\]</p>',
)


def append_diagram_tags(
    html: str,
    assets: Iterable["DiagramAsset"],
) -> str:
    """在正文末尾补插 LLM 漏写的图示 figure 标记(按 diagram_id 去重)。"""

    existing_ids = set(DIAGRAM_FIGURE_RE.findall(html or ""))
    parts: list[str] = []
    for asset in assets:
        if str(asset.diagram_id) in existing_ids:
            continue
        parts.append(
            f'<figure data-diagram-id="{asset.diagram_id}">'
            f"<figcaption>{asset.title}</figcaption></figure>"
        )
    if not parts:
        return html
    logger.info("图示漏插兜底追加: count={}", len(parts))
    return (html or "") + "\n" + "\n".join(parts)


def rewrite_diagram_figures(html: str) -> str:
    """把 figure 协议重写为渲染期可识别的占位段落(与素材图片占位同构)。"""

    def _replace(match: re.Match[str]) -> str:
        diagram_id = match.group(1)
        return (
            f'<p class="__bid_diagram__" data-diagram-id="{diagram_id}">'
            f"{DIAGRAM_PLACEHOLDER_TEXT}</p>"
        )

    return DIAGRAM_FIGURE_RE.sub(_replace, html or "")


def extract_diagram_captions(original_html: str) -> dict[int, str]:
    """从原始草稿 HTML 中提取 diagram_id → 图题映射。"""

    caption_re = re.compile(
        r'<figure\s+data-diagram-id="(\d+)"[^>]*>.*?'
        r"<figcaption>(.*?)</figcaption>",
        re.DOTALL,
    )
    return {
        int(match.group(1)): match.group(2).strip()
        for match in caption_re.finditer(original_html or "")
    }


def _ordered_diagram_refs(original_html: str) -> list[tuple[int, str]]:
    """按出现顺序提取 (diagram_id, caption); 与占位段落顺序一一对应。"""

    caption_re = re.compile(
        r'<figure\s+data-diagram-id="(\d+)"[^>]*>.*?'
        r"<figcaption>(.*?)</figcaption>",
        re.DOTALL,
    )
    ordered: dict[int, str] = {}
    for match in caption_re.finditer(original_html or ""):
        ordered[int(match.group(1))] = match.group(2).strip()
    return list(ordered.items())


_DIAGRAM_WIDTH_MM = {
    "flowchart": 130,
    "gantt": 150,
    "swimlane": 160,
    "roadmap": 160,
    "org_chart": 140,
    "layered": 140,
    "network": 160,
    "wbs": 150,
    "zone_layout": 150,
}
_PENDING_LABEL = "[图示待人工确认"


def _diagram_width_mm(record: object) -> int:
    diagram_type = str(getattr(record, "diagram_type", "") or "")
    return _DIAGRAM_WIDTH_MM.get(diagram_type, 130)


def _fitted_diagram_size_mm(record: object, png_path: str) -> tuple[float, float]:
    from PIL import Image

    from app.core.writer.visuals.page_fit import fit_diagram_size_mm

    width_mm = float(_diagram_width_mm(record))
    with Image.open(png_path) as image:
        px_w, px_h = image.size
    return fit_diagram_size_mm(width_mm, px_w, px_h)


def insert_diagram_figures(
    doc: object,
    original_html: str,
    Pt: type,
    diagram_records: dict[int, object],
    start_para_index: int = 0,
) -> None:
    """把 Word 文档中的图示占位段落替换为已渲染 PNG(顺序与原文 figure 一致)。"""

    from pathlib import Path as _Path

    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm
    from docx.text.paragraph import Paragraph

    refs = _ordered_diagram_refs(original_html)
    if not refs:
        return

    cursor = 0
    for paragraph in doc.paragraphs[start_para_index:]:
        if DIAGRAM_PLACEHOLDER_TEXT not in paragraph.text:
            continue
        if cursor >= len(refs):
            break
        diagram_id, caption = refs[cursor]
        caption = sanitize_xml_compatible_text(caption)
        cursor += 1
        record = diagram_records.get(diagram_id)
        if record is None or getattr(record, "status", "ready") != "ready":
            logger.warning(
                "图示记录缺失或未就绪: diagram_id={}, caption={!r}",
                diagram_id, caption,
            )
            paragraph.text = f"{_PENDING_LABEL}: {caption}]"
            continue
        png_path = str(getattr(record, "png_path", "") or "")
        if not png_path or not _Path(png_path).exists():
            logger.warning(
                "图示文件缺失: diagram_id={}, path={!r}", diagram_id, png_path,
            )
            paragraph.text = f"{_PENDING_LABEL}: {caption}]"
            continue

        paragraph.clear()
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        try:
            fitted_w, _fitted_h = _fitted_diagram_size_mm(record, png_path)
            run.add_picture(png_path, width=Mm(fitted_w))
        except Exception as exc:  # noqa: BLE001 - 单图失败仅标记
            logger.warning("图示插入失败: diagram_id={}, err={}", diagram_id, exc)
            paragraph.text = f"[图示插入失败: {caption}]"
            continue

        caption_text = sanitize_xml_compatible_text(
            caption or str(getattr(record, "title", "")),
        )
        cap_element = paragraph._element.makeelement(
            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p", {},
        )
        paragraph._element.addnext(cap_element)
        cap_paragraph = Paragraph(cap_element, paragraph._parent)
        cap_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_paragraph.paragraph_format.first_line_indent = Pt(0)
        cap_run = cap_paragraph.add_run(caption_text)
        cap_run.font.size = Pt(9)


def mark_residual_diagram_placeholders(doc: object) -> None:
    """全文残留图示占位统一标记为待人工确认(不做全局补位)。"""

    marked = 0
    for paragraph in doc.paragraphs:
        if DIAGRAM_PLACEHOLDER_TEXT not in paragraph.text:
            continue
        paragraph.text = f"{_PENDING_LABEL}]"
        marked += 1
    if marked:
        logger.info("图示残留检测: {}处占位符标记为待人工确认", marked)


_PREVIEW_PLACEHOLDER_RE = re.compile(r"<p[^>]*>\[生成图示占位\]</p>")


def inline_diagram_figures(
    rendered_html: str,
    raw_html: str,
    diagram_records: dict[int, object],
) -> str:
    """预览 HTML 中的图示占位替换为 base64 内联图(与 Word 判定对齐)。"""

    import base64
    from html import escape
    from pathlib import Path as _Path

    refs = _ordered_diagram_refs(raw_html)
    if not refs:
        return rendered_html

    cursor = {"index": 0}

    def _replace(_match: re.Match[str]) -> str:
        index = cursor["index"]
        cursor["index"] = index + 1
        if index >= len(refs):
            return f'<p class="placeholder">{_PENDING_LABEL}]</p>'
        diagram_id, caption = refs[index]
        safe_caption = escape(caption or "")
        record = diagram_records.get(diagram_id)
        if record is None or getattr(record, "status", "ready") != "ready":
            return f'<p class="placeholder">{_PENDING_LABEL}: {safe_caption}]</p>'
        png_path = str(getattr(record, "png_path", "") or "")
        path = _Path(png_path) if png_path else None
        if path is None or not path.exists():
            return f'<p class="placeholder">[图示缺失: {safe_caption}]</p>'
        try:
            encoded = base64.b64encode(path.read_bytes()).decode("ascii")
        except OSError as exc:
            logger.warning("预览图示读取失败: path={}, err={}", png_path, exc)
            return f'<p class="placeholder">[图示缺失: {safe_caption}]</p>'
        fitted_w, fitted_h = _fitted_diagram_size_mm(record, png_path)
        return (
            '<figure class="bid-figure" style="text-align:center;margin:12pt 0;">'
            f'<img src="data:image/png;base64,{encoded}" alt="{safe_caption}" '
            f'style="width:{fitted_w:.1f}mm;max-width:100%;'
            f'max-height:{fitted_h:.1f}mm;height:auto;" />'
            f'<figcaption style="font-size:9pt;color:#333;margin-top:4pt;">'
            f"{safe_caption}</figcaption>"
            "</figure>"
        )

    return _PREVIEW_PLACEHOLDER_RE.sub(_replace, rendered_html)


__all__ = [
    "DIAGRAM_FIGURE_RE",
    "DIAGRAM_PLACEHOLDER_TEXT",
    "append_diagram_tags",
    "extract_diagram_captions",
    "inline_diagram_figures",
    "insert_diagram_figures",
    "mark_residual_diagram_placeholders",
    "rewrite_diagram_figures",
]

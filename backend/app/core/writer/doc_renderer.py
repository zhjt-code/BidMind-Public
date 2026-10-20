"""Word 文档渲染引擎 + HTML 预览生成

DocRenderer:  将目录树+草稿渲染为 Word (.docx) 字节流
HtmlPreviewBuilder: 将目录树+草稿组装为 HTML 预览字符串

HTML -> docx 转换使用 htmldocx 库(基于 BeautifulSoup 解析)。
渲染完成后通过 _enforce_format_baseline 强制扫描全文,
确保字体/颜色/缩进/斜体等格式符合标书规范基线。
"""

import re
import zipfile
from html import escape
from io import BytesIO
from pathlib import Path
from typing import Optional

from loguru import logger
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.writer.bid_format_defaults import (
    BODY_FIRST_INDENT_PT,
    BODY_SIZE_PT,
    COLOR_BLACK,
    FONT_BODY,
    FONT_HEADER_FOOTER,
    FONT_HEADING,
    HEADER_FOOTER_SIZE_PT,
    HEADING_SPECS,
    MARGIN_BOTTOM_MM,
    MARGIN_LEFT_MM,
    MARGIN_RIGHT_MM,
    MARGIN_TOP_MM,
    PAGE_HEIGHT_MM,
    PAGE_WIDTH_MM,
    TABLE_BODY_SIZE_PT,
    TABLE_HEADER_BOLD,
    TABLE_HEADER_SIZE_PT,
    merge_with_profile,
    normalize_line_spacing,
)
from app.core.writer.binding_consumption_policy import (
    select_confirmed_material_bindings,
)
from app.core.writer.numbering import (
    compute_chapter_numbers,
    format_number_for_depth,
    get_min_child_depth,
    strip_existing_number,
    strip_existing_number_strict,
)
from app.core.writer.section_composer import sanitize_residual_placeholders
from app.core.writer.section_structure import (
    SectionBlock,
    normalize_section_html,
    parse_section_blocks,
)
from app.core.writer.verbatim_section_renderer import (
    is_fixed_template_node,
    render_verbatim_preview_html,
    write_verbatim_docx_section,
)
from app.core.xml_compat import sanitize_xml_compatible_text
from app.models.writing_binding import BindingType, WritingBinding
from app.models.writing_material_image import WritingMaterialImage
from app.models.writing_outline_node import WritingOutlineNode
from app.models.writing_render_profile import WritingRenderProfile
from app.models.writing_section_draft import WritingSectionDraft

# ---------------------------------------------------------------------------
# fix: htmldocx 按首行列数创建 python-docx Table，后续行列数超出
# 即 IndexError。归一化策略：
#   1. 多组 thead/tbody -> 拆成多个独立 <table>
#   2. 每张表内行列数对齐到最大值（补空 <td>/<th>）
# ---------------------------------------------------------------------------

def _normalize_table_html_for_docx(html: str) -> str:
    """将可能触发 HtmlToDocx IndexError 的表格 HTML 归一化。

    - 拆分单表内多组 thead/tbody 为多张标准 table
    - 统一行列数，避免后续行列数 > 首行
    返回可安全交给 HtmlToDocx 的 HTML 字符串。
    """
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    tables = soup.find_all("table")
    if not tables:
        return html

    changed = False
    for table in list(tables):
        fragments = _split_multi_section_table(table)
        if fragments is not None:
            changed = True
            for frag in reversed(fragments):
                table.insert_after(frag)
            table.decompose()
            continue

        if _pad_table_columns(table):
            changed = True

    return str(soup) if changed else html


def _split_multi_section_table(table: object) -> "list | None":
    """若 table 含多组 thead/tbody，拆成多张独立 table 并返回；否则返回 None。"""
    from bs4 import Tag

    theads = table.find_all("thead", recursive=False)
    tbodies = table.find_all("tbody", recursive=False)
    if len(theads) <= 1 and len(tbodies) <= 1:
        return None

    groups: list[list[Tag]] = []
    current: list[Tag] = []
    for child in table.children:
        if not isinstance(child, Tag):
            continue
        if child.name == "thead" and current:
            groups.append(current)
            current = []
        if child.name in {"thead", "tbody", "tfoot", "tr", "colgroup", "caption"}:
            current.append(child)
    if current:
        groups.append(current)
    if len(groups) <= 1:
        return None

    from bs4 import BeautifulSoup
    result: list[Tag] = []
    for group in groups:
        new_table = BeautifulSoup("<table></table>", "html.parser").table
        for tag in group:
            tag.extract()
            new_table.append(tag)
        _pad_table_columns(new_table)
        result.append(new_table)
    return result


def _pad_table_columns(table: object) -> bool:
    """统一表内所有行到相同列数（不足的尾部补空 <td>/<th>）。返回是否有修改。"""
    from bs4 import BeautifulSoup, Tag

    rows = table.find_all("tr", recursive=True)
    if not rows:
        return False

    def _cell_count(row: Tag) -> int:
        return len(row.find_all(["td", "th"], recursive=False))

    max_cols = max(_cell_count(r) for r in rows)
    if max_cols == 0:
        return False

    padded = False
    for row in rows:
        deficit = max_cols - _cell_count(row)
        if deficit > 0:
            is_header = row.find_parent("thead") is not None
            cell_tag = "th" if is_header else "td"
            for _ in range(deficit):
                new_cell = BeautifulSoup(
                    f"<{cell_tag}></{cell_tag}>", "html.parser",
                ).find(cell_tag)
                row.append(new_cell)
            padded = True
    return padded

_DOC_STYLE_BY_DEPTH = {
    0: "BM_H1",
    1: "BM_H2",
    2: "BM_H3",
    3: "BM_H4",
    4: "BM_H5",
}
_HTML_TAG_BY_DEPTH = {
    0: "h1",
    1: "h2",
    2: "h3",
    3: "h4",
    4: "h5",
}
def _set_cjk_font(rpr_or_style_element, font_name: str, font_name_ascii: str = ""):
    """设置 run 或 style 元素的中文+西文字体

    python-docx 的 font.name 只设置 w:ascii / w:hAnsi,
    中文/CJK 文字需要额外设置 w:eastAsia 属性。
    """
    from docx.oxml.ns import qn

    ascii_name = font_name_ascii or font_name
    rfonts = rpr_or_style_element.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr_or_style_element.insert(0, rfonts)

    rfonts.set(qn("w:ascii"), ascii_name)
    rfonts.set(qn("w:hAnsi"), ascii_name)
    rfonts.set(qn("w:eastAsia"), font_name)
    # 清除 theme 属性,防止 theme 覆盖显式字体
    for theme_attr in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
        key = qn(f"w:{theme_attr}")
        if key in rfonts.attrib:
            del rfonts.attrib[key]


def _set_font_color(rpr_element, hex_color: str):
    """设置 run properties 的字体颜色"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    color_el = rpr_element.find(qn("w:color"))
    if color_el is None:
        color_el = OxmlElement("w:color")
        rpr_element.append(color_el)
    color_el.set(qn("w:val"), hex_color)
    # 清除 theme 色,防止覆盖
    for attr in ["themeColor", "themeTint", "themeShade"]:
        key = qn(f"w:{attr}")
        if key in color_el.attrib:
            del color_el.attrib[key]


def _remove_italic(rpr_element):
    """确保 rPr 中没有斜体标记"""
    from docx.oxml.ns import qn

    for tag in ["w:i", "w:iCs"]:
        el = rpr_element.find(qn(tag))
        if el is not None:
            rpr_element.remove(el)


def _set_outline_level(style_element, level: int) -> None:
    """为段落样式写入 Word 大纲级别（0=一级，1=二级，2=三级）。

    职责: 让自定义标题样式可被 Word/WPS「自动目录」按大纲级别收集。
    不做: 不插入 TOC 域、不计算页码、不改标题文案。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pPr = style_element.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        style_element.append(pPr)

    outline = pPr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        pPr.append(outline)
    outline.set(qn("w:val"), str(level))


def _clear_outline_level(style_element) -> None:
    """移除段落样式上的大纲级别，避免四级及以下标题进入自动目录。"""
    from docx.oxml.ns import qn

    pPr = style_element.find(qn("w:pPr"))
    if pPr is None:
        return
    outline = pPr.find(qn("w:outlineLvl"))
    if outline is not None:
        pPr.remove(outline)


def _reset_list_numbering(doc: object, qn, OxmlElement) -> None:
    """为每组连续 List Number 段落重置编号从 1 开始

    htmldocx 使用 python-docx 的 "List Number" 样式，该样式共享
    全局编号序列。本函数找到每组连续的 List Number 段落，为每组
    的首段添加 numbering restart 标记。
    """
    numbering_part = doc.part.numbering_part
    if numbering_part is None:
        return

    num_xml = numbering_part._element
    existing_abstract_ids = set()
    for abstract_num in num_xml.findall(qn("w:abstractNum")):
        aid = abstract_num.get(qn("w:abstractNumId"))
        if aid:
            existing_abstract_ids.add(int(aid))

    next_abstract_id = max(existing_abstract_ids, default=0) + 100
    next_num_id = 1000

    prev_was_list_number = False
    current_num_id = None

    for paragraph in doc.paragraphs:
        is_list_number = (paragraph.style
                          and paragraph.style.name == "List Number")

        if is_list_number:
            if not prev_was_list_number:
                # 新组: 创建新的 numbering 定义
                abstract_id = next_abstract_id
                current_num_id = next_num_id
                next_abstract_id += 1
                next_num_id += 1

                abstract_num = OxmlElement("w:abstractNum")
                abstract_num.set(qn("w:abstractNumId"), str(abstract_id))

                lvl = OxmlElement("w:lvl")
                lvl.set(qn("w:ilvl"), "0")
                start = OxmlElement("w:start")
                start.set(qn("w:val"), "1")
                lvl.append(start)
                num_fmt = OxmlElement("w:numFmt")
                num_fmt.set(qn("w:val"), "decimal")
                lvl.append(num_fmt)
                lvl_text = OxmlElement("w:lvlText")
                lvl_text.set(qn("w:val"), "%1.")
                lvl.append(lvl_text)
                lvl_jc = OxmlElement("w:lvlJc")
                lvl_jc.set(qn("w:val"), "left")
                lvl.append(lvl_jc)
                abstract_num.append(lvl)
                num_xml.append(abstract_num)

                num_el = OxmlElement("w:num")
                num_el.set(qn("w:numId"), str(current_num_id))
                abstract_ref = OxmlElement("w:abstractNumId")
                abstract_ref.set(qn("w:val"), str(abstract_id))
                num_el.append(abstract_ref)
                num_xml.append(num_el)

            # 组内所有段落都设置相同的 numPr
            pPr = paragraph._element.find(qn("w:pPr"))
            if pPr is None:
                pPr = OxmlElement("w:pPr")
                paragraph._element.insert(0, pPr)

            old_numPr = pPr.find(qn("w:numPr"))
            if old_numPr is not None:
                pPr.remove(old_numPr)

            numPr = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl")
            ilvl.set(qn("w:val"), "0")
            numPr.append(ilvl)
            numId_el = OxmlElement("w:numId")
            numId_el.set(qn("w:val"), str(current_num_id))
            numPr.append(numId_el)
            pPr.append(numPr)

        prev_was_list_number = is_list_number


def _fix_zip_for_word(raw: bytes) -> bytes:
    """用全新 ZipInfo 重写 ZIP, 消除 Linux 环境残留的不兼容元数据.

    python-docx 在 Linux/Docker 上生成的 docx 携带 Unix 系统标记
    (create_system=3) 及其他 Python zipfile API 不可见的平台特定字段,
    部分版本 Microsoft Word 会因此拒绝打开. 此函数以全新 ZipInfo
    重新打包, 仅保留文件名和内容, 不改变任何 XML 或媒体数据.
    """
    src = BytesIO(raw)
    dst = BytesIO()
    with zipfile.ZipFile(src) as zin:
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for old_info in zin.infolist():
                new_info = zipfile.ZipInfo(old_info.filename)
                new_info.compress_type = zipfile.ZIP_DEFLATED
                new_info.create_system = 0
                zout.writestr(new_info, zin.read(old_info.filename))
    return dst.getvalue()


class DocRenderer:
    """Word 文档渲染器"""

    def __init__(self, db: AsyncSession) -> None:
        self.db = db

    async def render(self, project_id: int) -> bytes:
        """渲染完整 Word 文档"""
        try:
            from docx import Document as DocxDocument
            from docx.shared import Mm, Pt
        except ImportError:
            logger.error("python-docx 未安装，无法渲染 Word 文档")
            raise RuntimeError("python-docx 未安装")

        logger.info(f"Word渲染启动: project={project_id}")

        from app.core.writer.render_profile_currentness import ensure_render_profile_current
        await ensure_render_profile_current(self.db, project_id)
        profile = await self._load_render_profile(project_id)
        nodes = await self._load_outline_tree(project_id)
        drafts = await self._load_drafts_map(project_id)

        flat_nodes = self._flatten_tree(nodes)
        number_map = compute_chapter_numbers(nodes)
        image_records = await self._load_all_image_records(project_id)
        node_image_map = await self._build_node_image_map(project_id, image_records)
        consumable_image_records = {
            rec_id: rec
            for node_images in node_image_map.values()
            for rec_id, rec in node_images.items()
        }
        # P4: 生成图示为渲染增强项, 加载失败降级为占位标记
        diagram_records: dict[int, object] = {}
        try:
            from app.core.writer.visuals.repository import SectionDiagramRepository
            diagram_records = await SectionDiagramRepository(self.db).get_ready_by_project(
                project_id,
            )
        except Exception as exc:
            logger.warning(f"生成图示加载失败,降级为占位: project={project_id}, err={exc}")
        cfg = merge_with_profile(profile)

        drafts_with_content = sum(1 for d in drafts.values() if d.content)
        logger.info(
            f"渲染数据加载: nodes={len(flat_nodes)}, "
            f"drafts={len(drafts)}(有内容={drafts_with_content}), "
            f"images={len(image_records)}, "
            f"render_profile={'自定义' if profile else '默认'}"
        )

        doc = DocxDocument()
        consumed_image_keys: dict[str, str] = {}

        self._apply_styles(doc, cfg, Pt, Mm)

        empty_sections = 0
        for node in flat_nodes:
            draft = drafts.get(node.id)
            if not draft or not draft.content:
                empty_sections += 1
            self._write_section(
                doc, node, draft, Pt,
                number_map, node_image_map.get(node.id, {}), consumed_image_keys,
                diagram_records,
            )

        if empty_sections > 0:
            logger.warning(f"渲染警告: {empty_sections}/{len(flat_nodes)} 个章节无内容")

        self._set_header_footer(doc, cfg, Pt)
        self._replace_all_image_placeholders(doc, consumable_image_records, Pt)
        from app.core.writer.visuals.figure_bridge import (
            mark_residual_diagram_placeholders,
        )
        mark_residual_diagram_placeholders(doc)
        self._enforce_format_baseline(doc, cfg, Pt)
        self._verify_render_output(doc, consumable_image_records)

        buffer = BytesIO()
        doc.save(buffer)
        result = _fix_zip_for_word(buffer.getvalue())
        logger.info(
            f"Word渲染完成: project={project_id}, "
            f"size={len(result)/1024:.1f}KB, sections={len(flat_nodes)}"
        )
        return result

    # ==================== 样式 ====================

    def _apply_styles(
        self, doc: object, cfg: dict, Pt: type, Mm: type,
    ) -> None:
        """应用全局样式: 页面设置 + 正文字体 + 项目专用标题样式"""
        from docx.enum.style import WD_STYLE_TYPE
        from docx.oxml.ns import qn

        cfg["line_spacing"] = normalize_line_spacing(cfg.get("line_spacing"))

        # -- 页面设置 (A4 + 边距) --
        for section in doc.sections:
            section.page_width = Mm(PAGE_WIDTH_MM)
            section.page_height = Mm(PAGE_HEIGHT_MM)
            section.top_margin = Mm(cfg["margin_top"])
            section.bottom_margin = Mm(cfg["margin_bottom"])
            section.left_margin = Mm(cfg["margin_left"])
            section.right_margin = Mm(cfg["margin_right"])

        # -- Normal 样式: 正文字体(CJK + Latin) --
        normal = doc.styles["Normal"]
        normal.font.size = Pt(cfg["font_size_body"])
        normal.font.bold = False
        normal.font.italic = False
        normal.paragraph_format.line_spacing = cfg["line_spacing"]

        normal_rpr = normal.element.find(qn("w:rPr"))
        if normal_rpr is None:
            from docx.oxml import OxmlElement
            normal_rpr = OxmlElement("w:rPr")
            normal.element.append(normal_rpr)
        _set_cjk_font(normal_rpr, cfg["font_body"])
        _set_font_color(normal_rpr, COLOR_BLACK)
        _remove_italic(normal_rpr)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.first_line_indent = Pt(BODY_FIRST_INDENT_PT)

        if "BM_BODY" not in doc.styles:
            body_style = doc.styles.add_style("BM_BODY", WD_STYLE_TYPE.PARAGRAPH)
        else:
            body_style = doc.styles["BM_BODY"]
        body_style.base_style = normal
        body_style.font.size = Pt(cfg["font_size_body"])
        body_style.font.bold = False
        body_style.font.italic = False
        body_style.paragraph_format.line_spacing = cfg["line_spacing"]
        body_style.paragraph_format.space_before = Pt(0)
        body_style.paragraph_format.space_after = Pt(0)
        body_style.paragraph_format.first_line_indent = Pt(BODY_FIRST_INDENT_PT)
        body_rpr = body_style.element.find(qn("w:rPr"))
        if body_rpr is None:
            from docx.oxml import OxmlElement
            body_rpr = OxmlElement("w:rPr")
            body_style.element.append(body_rpr)
        _set_cjk_font(body_rpr, cfg["font_body"])
        _set_font_color(body_rpr, COLOR_BLACK)
        _remove_italic(body_rpr)

        # -- 项目标题样式: BM_H1 ~ BM_H5 --
        heading_sizes = {
            "BM_H1": cfg["font_size_h1"],
            "BM_H2": cfg["font_size_h2"],
            "BM_H3": cfg["font_size_h3"],
            "BM_H4": HEADING_SPECS["Heading 4"].size_pt,
            "BM_H5": HEADING_SPECS["Heading 5"].size_pt,
        }
        space_before_map = {
            "BM_H1": 24,
            "BM_H2": 18,
            "BM_H3": 14,
            "BM_H4": 10,
            "BM_H5": 6,
        }
        space_after_map = {
            "BM_H1": 8,
            "BM_H2": 6,
            "BM_H3": 4,
            "BM_H4": 2,
            "BM_H5": 0,
        }
        spec_key_map = {
            "BM_H1": "Heading 1",
            "BM_H2": "Heading 2",
            "BM_H3": "Heading 3",
            "BM_H4": "Heading 4",
            "BM_H5": "Heading 5",
        }
        for style_name, spec_key in spec_key_map.items():
            spec = HEADING_SPECS[spec_key]
            if style_name not in doc.styles:
                hs = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
            else:
                hs = doc.styles[style_name]
            hs.base_style = normal
            size = heading_sizes.get(style_name, spec.size_pt)
            hs.font.size = Pt(size)
            hs.font.bold = spec.bold
            hs.font.italic = False
            hs.paragraph_format.line_spacing = cfg["line_spacing"]
            hs.paragraph_format.space_before = Pt(space_before_map[style_name])
            hs.paragraph_format.space_after = Pt(space_after_map[style_name])
            hs.paragraph_format.first_line_indent = Pt(0)
            hs.paragraph_format.keep_with_next = True

            hs_rpr = hs.element.find(qn("w:rPr"))
            if hs_rpr is None:
                from docx.oxml import OxmlElement
                hs_rpr = OxmlElement("w:rPr")
                hs.element.append(hs_rpr)

            override_font = cfg["font_heading"] if style_name in ("BM_H1", "BM_H2") else spec.font_name
            _set_cjk_font(hs_rpr, override_font, spec.font_name_ascii)
            _set_font_color(hs_rpr, COLOR_BLACK)
            _remove_italic(hs_rpr)

            # Word「自动目录」按大纲级别收集；仅开放前三级
            outline_level = {"BM_H1": 0, "BM_H2": 1, "BM_H3": 2}.get(style_name)
            if outline_level is not None:
                _set_outline_level(hs.element, outline_level)
            else:
                _clear_outline_level(hs.element)

        if "BM_LEADIN" not in doc.styles:
            leadin_style = doc.styles.add_style("BM_LEADIN", WD_STYLE_TYPE.PARAGRAPH)
        else:
            leadin_style = doc.styles["BM_LEADIN"]
        leadin_style.base_style = body_style
        leadin_style.font.size = Pt(cfg["font_size_body"])
        leadin_style.font.bold = False
        leadin_style.font.italic = False
        leadin_style.paragraph_format.line_spacing = cfg["line_spacing"]
        leadin_style.paragraph_format.space_before = Pt(0)
        leadin_style.paragraph_format.space_after = Pt(0)
        leadin_style.paragraph_format.first_line_indent = Pt(0)

        # -- 兼容样式: 仅保留为兜底 --
        for list_style_name in ("List Bullet", "List Number"):
            if list_style_name in doc.styles:
                ls = doc.styles[list_style_name]
                ls.font.size = Pt(cfg["font_size_body"])
                ls.font.bold = False
                ls.font.italic = False
                ls_rpr = ls.element.find(qn("w:rPr"))
                if ls_rpr is not None:
                    _set_cjk_font(ls_rpr, cfg["font_body"])
                    _set_font_color(ls_rpr, COLOR_BLACK)
                    _remove_italic(ls_rpr)

    # ==================== 章节写入 ====================

    def _write_section(
        self, doc: object, node: WritingOutlineNode,
        draft: Optional[WritingSectionDraft],
        Pt: type,
        number_map: Optional[dict[int, str]] = None,
        image_records: Optional[dict[int, "WritingMaterialImage"]] = None,
        consumed_image_keys: Optional[dict[str, str]] = None,
        diagram_records: Optional[dict[int, object]] = None,
    ) -> None:
        """写入单个章节"""
        number = (number_map or {}).get(node.id, "")
        clean_title = strip_existing_number(node.title)
        display_title = sanitize_xml_compatible_text(
            f"{number}{clean_title}" if number else clean_title,
        )
        heading_para = doc.add_paragraph(style=self._get_doc_style_name(node.depth))
        heading_para.add_run(display_title)

        if not draft or not draft.content:
            doc.add_paragraph(
                sanitize_xml_compatible_text(f"[{node.title} - 内容待补充]"),
                style="BM_BODY",
            )
            return

        section_images = image_records
        if image_records:
            from app.core.writer.image_node_affinity import filter_images_for_node
            filtered = filter_images_for_node(
                list(image_records.values()), node.title,
            )
            section_images = {img.id: img for img in filtered} if filtered else {}

        if is_fixed_template_node(node):
            # 固定模板章节保真输出, 不走统一编号契约
            write_verbatim_docx_section(
                self, doc, draft.content, Pt, section_images, clean_title,
                consumed_image_keys,
            )
            return

        self._html_to_docx(
            doc, draft.content, Pt, section_images, clean_title, node.depth,
            consumed_image_keys,
            diagram_records,
        )

    def _html_to_docx(
        self, doc: object, html: str, Pt: type,
        image_records: Optional[dict[int, "WritingMaterialImage"]] = None,
        section_title: str = "",
        base_depth: int = 0,
        consumed_image_keys: Optional[dict[str, str]] = None,
        diagram_records: Optional[dict[int, object]] = None,
    ) -> None:
        """HTML -> python-docx 转换（显式编号 + HTML表格兜底）"""
        raw_html = html.strip()
        html = self._preprocess_html_for_bid(raw_html, section_title)
        blocks = self._merge_bare_numbering_with_following_paragraph(
            parse_section_blocks(html),
        )
        start_para_index = len(doc.paragraphs)
        heading_counters: dict[int, int] = {}
        context_depth = base_depth

        for block in blocks:
            if block.kind == "heading":
                context_depth = self._render_heading_block(
                    doc, block, base_depth, heading_counters, context_depth,
                )
            elif block.kind == "numbered_paragraph":
                context_depth = self._render_numbered_paragraph_block(
                    doc, block, base_depth, heading_counters, context_depth,
                )
            elif block.kind == "paragraph":
                self._render_body_paragraph(doc, block.text)
            elif block.kind == "list":
                self._render_list_block(doc, block, context_depth)
            elif block.kind in {"table", "html"}:
                self._append_html_fragment_to_doc(block.html, doc)

        self._insert_figure_images(
            doc,
            raw_html,
            Pt,
            image_records or {},
            start_para_index,
            section_title,
            consumed_image_keys,
        )
        # P4: 生成图示插入(独立于素材图片, 无亲和/独占闸门)
        from app.core.writer.visuals.figure_bridge import insert_diagram_figures
        insert_diagram_figures(
            doc,
            raw_html,
            Pt,
            diagram_records or {},
            start_para_index,
        )

    @staticmethod
    def _merge_bare_numbering_with_following_paragraph(
        blocks: list[SectionBlock],
    ) -> list[SectionBlock]:
        """将纯编号块与紧跟的正文段落合并。

        当一个 pure_numbering 块（如独立的 "（三）"）后面紧跟普通正文段落时，
        合并为一个 numbered_paragraph，使编号前缀和正文内容渲染在同一段落中。
        如果后面是标题、列表、表格等结构块，则保持独占一行（充当子节点分隔符）。
        """
        if not blocks:
            return blocks

        result: list[SectionBlock] = []
        skip_next = False

        for i, block in enumerate(blocks):
            if skip_next:
                skip_next = False
                continue

            if block.pure_numbering and i + 1 < len(blocks):
                next_block = blocks[i + 1]
                if next_block.kind == "paragraph" and next_block.text.strip():
                    result.append(SectionBlock(
                        kind="numbered_paragraph",
                        text=next_block.text,
                        explicit_depth=block.explicit_depth,
                        pure_numbering=False,
                        prefer_structure_depth=block.explicit_depth is not None,
                    ))
                    skip_next = True
                    continue

            result.append(block)

        return result

    @staticmethod
    def _preprocess_html_for_bid(html: str, section_title: str = "") -> str:
        """标书格式预处理

 0. 占位符兜底闸门 B1: 历史脏 draft 与 LLM 复读的 {xxx} 转 [待补充:xxx]
 1. 将 <figure> 替换为占位段落供后续图片插入
 2. 清理行首/行尾孤立 <br>
 3. 去除 HTML 中与章节标题重复的 <h> 标签
 4. 去除纯文本行中与章节标题重复的行
 (覆盖 fixed_template raw_text 中的纯文本标题)
 """
        html = sanitize_xml_compatible_text(html)
        html = sanitize_residual_placeholders(html)
        html = re.sub(
            r'<figure\s+data-image-id="(\d+)"[^>]*>.*?</figure>',
            r'<p class="__bid_figure__" data-image-id="\1">[图片占位]</p>',
            html,
            flags=re.DOTALL,
        )
        # P4: 生成图示 figure → 占位段落(由 visuals/figure_bridge 统一维护协议)
        from app.core.writer.visuals.figure_bridge import rewrite_diagram_figures
        html = rewrite_diagram_figures(html)
        html = re.sub(r'^(\s*<br\s*/?>)+', '', html)
        html = re.sub(r'(<br\s*/?>)+\s*$', '', html)
        return normalize_section_html(html, section_title)

    @staticmethod
    def _get_doc_style_name(depth: int) -> str:
        return _DOC_STYLE_BY_DEPTH.get(depth, "BM_H5")

    @staticmethod
    def _get_html_tag_name(depth: int) -> str:
        return _HTML_TAG_BY_DEPTH.get(depth, "h5")

    def _render_heading_block(
        self,
        doc: object,
        block: SectionBlock,
        base_depth: int,
        heading_counters: dict[int, int],
        current_depth: int,
    ) -> int:
        absolute_depth = self._resolve_block_depth(
            block,
            preferred_depth=min(
                base_depth + max(block.level, 1),
                max(current_depth + 1, max(base_depth + 1, 1)),
            ),
            base_depth=base_depth,
            source="heading",
        )
        if block.pure_numbering:
            paragraph = doc.add_paragraph(style=self._get_doc_style_name(absolute_depth))
            paragraph.add_run(self._next_number_prefix(absolute_depth, heading_counters))
            return absolute_depth
        clean_text = self._sanitize_numbered_text(block.text, "heading")
        if not clean_text:
            return absolute_depth
        display_text = self._build_numbered_text(
            clean_text, absolute_depth, heading_counters,
        )
        paragraph = doc.add_paragraph(style=self._get_doc_style_name(absolute_depth))
        paragraph.add_run(display_text)
        return absolute_depth

    def _render_numbered_paragraph_block(
        self,
        doc: object,
        block: SectionBlock,
        base_depth: int,
        heading_counters: dict[int, int],
        current_depth: int,
    ) -> int:
        absolute_depth = self._resolve_block_depth(
            block,
            preferred_depth=self._compute_numbered_preferred_depth(
                block,
                base_depth,
                current_depth,
            ),
            base_depth=base_depth,
            source="numbered_paragraph",
        )
        if block.pure_numbering:
            paragraph = doc.add_paragraph(style="BM_LEADIN")
            paragraph.add_run(self._next_number_prefix(absolute_depth, heading_counters))
            return absolute_depth
        clean_text = self._sanitize_numbered_text(block.text, "numbered_paragraph")
        if not clean_text:
            return absolute_depth
        display_text = self._build_numbered_text(
            clean_text, absolute_depth, heading_counters,
        )
        paragraph = doc.add_paragraph(style="BM_LEADIN")
        paragraph.add_run(display_text)
        return absolute_depth

    @staticmethod
    def _compute_numbered_preferred_depth(
        block: SectionBlock,
        base_depth: int,
        current_depth: int,
    ) -> int:
        """在父约束下推导正文编号块的目标层级。

        - 章节节点 `base_depth` 给出正文的最小子层级
        - `current_depth` 记录上一个结构块落点
        - 可见编号只作为 hint，最多允许比当前上下文再深入一层
        """
        minimum_child_depth = get_min_child_depth(base_depth)
        normalized_current = max(current_depth, minimum_child_depth - 1)
        explicit_depth = block.explicit_depth

        if explicit_depth is None:
            return max(normalized_current, minimum_child_depth)
        if explicit_depth < minimum_child_depth:
            return minimum_child_depth
        if explicit_depth <= normalized_current:
            return explicit_depth
        return min(explicit_depth, normalized_current + 1)

    @staticmethod
    def _append_html_fragment_to_doc(html_fragment: str, doc: object) -> None:
        """归一化 + HtmlToDocx 转换，失败时降级为纯文本。

        所有 table/html 块的 Word 渲染统一走此入口，保证归一化和兜底逻辑一致。
        """
        from htmldocx import HtmlToDocx

        html_fragment = sanitize_xml_compatible_text(html_fragment)
        normalized = _normalize_table_html_for_docx(html_fragment)
        parser = HtmlToDocx()
        parser.table_style = "Table Grid"
        try:
            parser.add_html_to_document(normalized, doc)
        except Exception as e:
            logger.error(f"HTML片段转换失败,降级为纯文本: {e}")
            paragraph = doc.add_paragraph(style="Normal")
            paragraph.add_run(re.sub(r"<[^>]+>", "", html_fragment))

    @staticmethod
    def _render_body_paragraph(doc: object, text: str) -> None:
        clean = sanitize_xml_compatible_text(text).strip()
        if not clean:
            return
        paragraph = doc.add_paragraph(style="BM_BODY")
        paragraph.add_run(clean)

    def _render_list_block(
        self,
        doc: object,
        block: SectionBlock,
        base_depth: int,
        nesting_level: int = 1,
    ) -> None:
        for idx, item in enumerate(block.items):
            clean_text = self._sanitize_numbered_text(item.text, "list_item")
            if not clean_text:
                continue
            prefix = ""
            if block.ordered:
                prefix = format_number_for_depth(base_depth + nesting_level, idx)
            text = f"{prefix}{clean_text}".strip()
            paragraph = doc.add_paragraph(style="BM_LEADIN" if block.ordered else "BM_BODY")
            paragraph.add_run(text)
            for child in item.children:
                if child.kind == "list":
                    self._render_list_block(doc, child, base_depth, nesting_level + 1)
                elif child.kind == "paragraph":
                    self._render_body_paragraph(doc, child.text)
                elif child.kind == "numbered_paragraph":
                    nested_text = self._sanitize_numbered_text(
                        child.text, "nested_numbered_paragraph",
                    )
                    if not nested_text:
                        continue
                    nested_para = doc.add_paragraph(style="BM_LEADIN")
                    nested_base_depth = base_depth + nesting_level
                    nested_depth = self._resolve_block_depth(
                        child,
                        preferred_depth=self._compute_numbered_preferred_depth(
                            child,
                            nested_base_depth,
                            nested_base_depth,
                        ),
                        base_depth=nested_base_depth,
                        source="nested_numbered_paragraph",
                    )
                    nested_prefix = format_number_for_depth(
                        nested_depth,
                        0,
                    )
                    nested_para.add_run(f"{nested_prefix}{nested_text}")
                elif child.kind == "heading":
                    self._render_heading_block(
                        doc,
                        child,
                        base_depth,
                        {},
                        base_depth + nesting_level - 1,
                    )
                elif child.kind in {"table", "html"}:
                    self._append_html_fragment_to_doc(child.html, doc)

    @staticmethod
    def _next_number_prefix(
        absolute_depth: int,
        heading_counters: dict[int, int],
    ) -> str:
        heading_counters[absolute_depth] = heading_counters.get(absolute_depth, 0) + 1
        for key in list(heading_counters.keys()):
            if key > absolute_depth:
                del heading_counters[key]
        return format_number_for_depth(
            absolute_depth,
            heading_counters[absolute_depth] - 1,
        )

    @classmethod
    def _build_numbered_text(
        cls,
        text: str,
        absolute_depth: int,
        heading_counters: dict[int, int],
    ) -> str:
        prefix = cls._next_number_prefix(absolute_depth, heading_counters)
        return f"{prefix}{text}".strip()

    @staticmethod
    def _resolve_block_depth(
        block: SectionBlock,
        *,
        preferred_depth: int,
        base_depth: int,
        source: str,
    ) -> int:
        explicit_depth = block.explicit_depth
        if explicit_depth is None:
            return preferred_depth

        minimum_child_depth = max(base_depth + 1, 1)
        if block.pure_numbering:
            if minimum_child_depth <= explicit_depth <= preferred_depth:
                return explicit_depth
            logger.warning(
                "编号纯标记层级冲突已回退: source={}, explicit={}, preferred={}, base={}",
                source,
                explicit_depth,
                preferred_depth,
                base_depth,
            )
            return preferred_depth

        if getattr(block, "prefer_structure_depth", False):
            if explicit_depth < minimum_child_depth or explicit_depth > preferred_depth:
                logger.warning(
                    "编号层级冲突已回退: source={}, explicit={}, preferred={}, base={}, text={!r}",
                    source,
                    explicit_depth,
                    preferred_depth,
                    base_depth,
                    block.text[:80],
                )
                return preferred_depth
        return explicit_depth

    @staticmethod
    def _sanitize_numbered_text(text: str, source: str) -> str:
        clean = strip_existing_number_strict(text or "").strip()
        if re.match(r"^[)）]\s+\S+", clean):
            logger.warning(f"括号残留被清洗: source={source}, text={text!r}")
            clean = re.sub(r"^[)）]\s*", "", clean).strip()
        if re.match(r"^[.．、:：;；]+\s*", clean):
            logger.warning(f"标点残留被清洗: source={source}, text={text!r}")
            clean = re.sub(r"^[.．、:：;；]+\s*", "", clean).strip()
        if re.fullmatch(r"[¥￥、.．:：;；,，()]+", clean):
            logger.warning(f"模板符号残留被丢弃: source={source}, text={text!r}")
            return ""
        if re.fullmatch(r"\d+", clean):
            logger.warning(f"编号残留被丢弃: source={source}, text={text!r}")
            return ""
        if re.match(r"^\d+\s+\S+", clean):
            logger.warning(f"编号残留被清洗: source={source}, text={text!r}")
            clean = re.sub(r"^\d+[\t \u3000]+", "", clean).strip()
        if re.match(r"^[（(][一二三四五六七八九十\d]+[)）]\s+\d+\s+", clean):
            logger.warning(f"双编号残留被清洗: source={source}, text={text!r}")
            clean = re.sub(r"^[（(][一二三四五六七八九十\d]+[)）]\s+\d+[\t \u3000]+", "", clean).strip()
        return clean

    @staticmethod
    def _calc_image_width(rec: "WritingMaterialImage") -> object:
        """根据图片元数据计算插入宽度(保真策略)

        有尺寸: 按 150DPI 换算物理宽度,证照类至少 120mm
        无尺寸: 证照类 140mm,其他 120mm
        统一上限 160mm,下限 60mm
        """
        from docx.shared import Mm

        from app.models.writing_material_image import ImageType

        MAX_W = Mm(160)
        MIN_W = Mm(60)
        CERT_MIN = Mm(120)

        is_cert = rec.image_type in (
            ImageType.license, ImageType.personnel_cert,
            ImageType.qualification,
        ) if rec.image_type else False

        w_px = rec.width or 0
        if w_px > 0:
            physical_mm = w_px * 25.4 / 150
            target = Mm(physical_mm)
            if is_cert:
                target = max(target, CERT_MIN)
            target = max(min(target, MAX_W), MIN_W)
            return target

        return Mm(140) if is_cert else Mm(120)

    @staticmethod
    def _insert_figure_images(
        doc: object, original_html: str, Pt: type,
        image_records: dict[int, "WritingMaterialImage"],
        start_para_index: int = 0,
        current_section_title: str = "",
        consumed_image_keys: Optional[dict[str, str]] = None,
    ) -> None:
        """扫描文档中的图片占位段落,替换为实际图片"""
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        from app.core.writer.image_node_affinity import filter_images_for_node

        figure_pattern = re.compile(
            r'data-image-id="(\d+)"[^>]*>.*?'
            r'<figcaption>(.*?)</figcaption>',
            re.DOTALL,
        )
        image_map: dict[str, str] = {}
        for match in figure_pattern.finditer(original_html):
            image_map[match.group(1)] = match.group(2).strip()

        if not image_map:
            return

        identity_candidates: list[WritingMaterialImage] = []
        for img_id_str in image_map:
            rec = image_records.get(int(img_id_str))
            if rec is None:
                continue
            if current_section_title and not filter_images_for_node([rec], current_section_title):
                continue
            identity_candidates.append(rec)
        from app.core.writer.identity_image_layout import build_identity_image_layout
        identity_layout = build_identity_image_layout(identity_candidates)
        identity_ids = {str(image_id) for image_id in identity_layout.ordered_image_ids}
        identity_rendered = False
        image_order = list(image_map.keys())
        placeholder_index = 0

        for paragraph in list(doc.paragraphs[start_para_index:]):
            if "[图片占位]" not in paragraph.text:
                continue
            if placeholder_index >= len(image_order):
                paragraph.text = "[图片待人工确认]"
                continue

            img_id_str = image_order[placeholder_index]
            placeholder_index += 1
            caption = sanitize_xml_compatible_text(image_map.get(img_id_str, ""))

            if img_id_str in identity_ids and identity_rendered:
                paragraph.clear()
                parent = paragraph._element.getparent()
                if parent is not None:
                    parent.remove(paragraph._element)
                continue

            if img_id_str in identity_ids:
                from app.core.writer.identity_image_docx import (
                    insert_identity_image_block,
                )
                consumed_ids = insert_identity_image_block(
                    doc,
                    paragraph,
                    identity_layout,
                    {int(k): v for k, v in image_map.items()},
                    Pt,
                    consumed_image_keys,
                    current_section_title,
                )
                identity_rendered = True
                for consumed_id in consumed_ids:
                    image_map.pop(str(consumed_id), None)
                continue

            rec = image_records.get(int(img_id_str))
            if not rec:
                paragraph.text = f"[图片待人工确认: {caption}]"
                image_map.pop(img_id_str, None)
                continue
            if current_section_title and not filter_images_for_node([rec], current_section_title):
                logger.warning(
                    "图片最终闸门拦截: image_id={}, section={!r}, caption={!r}",
                    rec.id,
                    current_section_title,
                    caption,
                )
                paragraph.text = f"[图片待人工确认: {caption}]"
                image_map.pop(img_id_str, None)
                continue
            if not Path(rec.image_path).exists():
                paragraph.text = f"[图片缺失: {caption}]"
                image_map.pop(img_id_str, None)
                continue

            if consumed_image_keys is not None and DocRenderer._is_exclusive_image(rec):
                consume_key = rec.file_hash or rec.image_path
                prev_section = consumed_image_keys.get(consume_key)
                if prev_section and prev_section != current_section_title:
                    logger.warning(
                        "图片重复复用拦截: image_id={}, prev_section={!r}, current_section={!r}",
                        rec.id,
                        prev_section,
                        current_section_title,
                    )
                    paragraph.text = f"[图片待人工确认: {caption}]"
                    image_map.pop(img_id_str, None)
                    continue

            paragraph.clear()
            # 与 figure_bridge / 预览 text-align:center 对齐：取消正文首行缩进并居中，
            # 否则 BM_BODY 的 2 字符缩进会让图片视觉右偏。
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Pt(0)
            run = paragraph.add_run()
            try:
                width = DocRenderer._calc_image_width(rec)
                run.add_picture(rec.image_path, width=width)
            except Exception as e:
                logger.warning(f"图片插入失败: id={rec.id}, err={e}")
                paragraph.text = f"[图片插入失败: {caption}]"
                continue

            cap_para = paragraph._element.makeelement(
                '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p', {}
            )
            paragraph._element.addnext(cap_para)
            from docx.text.paragraph import Paragraph
            cap_p = Paragraph(cap_para, paragraph._parent)
            cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap_p.paragraph_format.first_line_indent = Pt(0)
            cap_run = cap_p.add_run(caption)
            cap_run.font.size = Pt(9)

            if consumed_image_keys is not None and DocRenderer._is_exclusive_image(rec):
                consume_key = rec.file_hash or rec.image_path
                consumed_image_keys[consume_key] = current_section_title

            del image_map[img_id_str]

    @staticmethod
    def _is_exclusive_image(rec: "WritingMaterialImage") -> bool:
        """转发到 image_node_affinity.is_exclusive_image(单一真源,与预览共用)。"""
        from app.core.writer.image_node_affinity import is_exclusive_image
        return is_exclusive_image(rec)

    # ==================== 图片残留占位符检测(不做全局补位) ====================

    @staticmethod
    def _replace_all_image_placeholders(
        doc: object,
        image_records: dict[int, "WritingMaterialImage"],
        Pt: type,
    ) -> None:
        """全文扫描残留 [图片占位] 段落,标记为待人工确认

        P3 改造: 不再从项目级剩余图片池中全局补位,
        避免上游匹配不准导致图片错插到不相关章节。
        所有图片插入应在 _insert_figure_images (精确ID匹配)中完成,
        残留的占位符统一标记为 [图片待人工确认] 供用户审核。
        """
        if not image_records:
            return

        current_section = ""
        marked = 0

        for paragraph in doc.paragraphs:
            style_name = paragraph.style.name if paragraph.style else ""
            if "Heading" in style_name or style_name.startswith("BM_H"):
                current_section = paragraph.text.strip()
                continue

            if "[图片占位]" not in paragraph.text:
                continue

            paragraph.text = (
                f"[图片待人工确认: 章节={current_section[:30]}]"
            )
            marked += 1

        if marked > 0:
            logger.info(
                f"图片残留检测: {marked}处占位符标记为待人工确认"
            )

    # ==================== 格式基线强制执行 ====================

    @staticmethod
    def _enforce_format_baseline(doc: object, cfg: dict, Pt: type) -> None:
        """全文扫描强制格式合规

        渲染完成后统一执行,作为最终闸门确保:
        - 全部段落行距统一(修正 htmldocx 硬编码的 1.0)
        - 全部 run 设置正确的 CJK 字体
        - 全部 run 颜色为黑色
        - 清除所有斜体
        - 正文段落首行缩进(排除 TOC 域代码段落)
        - 表格字体和大小
        - 清理多余空段落
        """
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        heading_styles = {"BM_H1", "BM_H2", "BM_H3", "BM_H4", "BM_H5"}
        list_styles = {"List Bullet", "List Number"}
        body_font = cfg.get("font_body", FONT_BODY)
        target_spacing = normalize_line_spacing(cfg.get("line_spacing"))

        enforced_runs = 0
        enforced_indent = 0
        enforced_spacing = 0
        removed_empty = 0

        for paragraph in doc.paragraphs:
            style_name = paragraph.style.name if paragraph.style else ""

            # P9: 统一行距 (htmldocx 对 list 硬编码 1.0)
            if (paragraph.paragraph_format.line_spacing is not None
                    and paragraph.paragraph_format.line_spacing != target_spacing):
                paragraph.paragraph_format.line_spacing = target_spacing
                enforced_spacing += 1

            # P11: 跳过含域代码的段落 (TOC)
            has_field = bool(paragraph._element.findall(f'.//{qn("w:fldChar")}'))

            # 正文段落首行缩进 (排除标题/列表/域代码段落)
            if (
                style_name not in heading_styles
                and style_name not in list_styles
                and paragraph.text.strip()
                and not has_field
                and paragraph.paragraph_format.first_line_indent is None
            ):
                paragraph.paragraph_format.first_line_indent = Pt(BODY_FIRST_INDENT_PT)
                enforced_indent += 1

            # 逐 run 扫描
            for run in paragraph.runs:
                rpr = run._r.find(qn("w:rPr"))
                if rpr is None:
                    rpr = OxmlElement("w:rPr")
                    run._r.insert(0, rpr)

                rfonts = rpr.find(qn("w:rFonts"))
                ea_attr = qn("w:eastAsia")
                if rfonts is None or ea_attr not in rfonts.attrib:
                    if style_name in heading_styles:
                        spec = HEADING_SPECS.get(style_name)
                        if spec:
                            _set_cjk_font(rpr, spec.font_name, spec.font_name_ascii)
                    else:
                        _set_cjk_font(rpr, body_font)
                    enforced_runs += 1

                _set_font_color(rpr, COLOR_BLACK)
                _remove_italic(rpr)

        # 表格格式
        for table in doc.tables:
            for ri, row in enumerate(table.rows):
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.first_line_indent = None
                        for run in paragraph.runs:
                            rpr = run._r.find(qn("w:rPr"))
                            if rpr is None:
                                rpr = OxmlElement("w:rPr")
                                run._r.insert(0, rpr)
                            _set_cjk_font(rpr, cfg.get("font_body", FONT_BODY))
                            _set_font_color(rpr, COLOR_BLACK)
                            _remove_italic(rpr)
                            run.font.size = Pt(TABLE_BODY_SIZE_PT)
                            if ri == 0:
                                run.font.bold = TABLE_HEADER_BOLD
                                run.font.size = Pt(TABLE_HEADER_SIZE_PT)

        # 兼容兜底: 若仍有 htmldocx 生成的 List Number，则尽量按组重置
        _reset_list_numbering(doc, qn, OxmlElement)

        # 清理空段落 (含文本/域代码/图片的段落保留，其余删除)
        body = doc.element.body
        w_p = qn("w:p")
        for p_el in list(body.findall(w_p)):
            texts = p_el.findall(f'.//{qn("w:t")}')
            fields = p_el.findall(f'.//{qn("w:fldChar")}')
            drawings = p_el.findall(f'.//{qn("w:drawing")}')
            has_text = any(t.text and t.text.strip() for t in texts)
            if has_text or fields or drawings:
                continue
            body.remove(p_el)
            removed_empty += 1

        logger.info(
            f"格式基线执行: runs={enforced_runs}, indents={enforced_indent}, "
            f"spacing={enforced_spacing}, empty_removed={removed_empty}"
        )

    # ==================== 渲染质检闸门 ====================

    @staticmethod
    def _verify_render_output(
        doc: object,
        image_records: dict[int, "WritingMaterialImage"],
    ) -> None:
        """渲染完成后的出厂质检 -- 只读检查，不修改文档

        error: 文档结构缺陷(残留占位、重复标题等)
        warning: 非阻断性质检信息(图片缺口、待确认占位等)
        新增检查项时只需追加代码，永远不删除已有检查。
        """
        from docx.oxml.ns import qn

        error_issues: list[str] = []
        warning_issues: list[str] = []

        # 检查1: [图片占位] 不应残留
        for p in doc.paragraphs:
            if "[图片占位]" in p.text:
                error_issues.append(f"残留 [图片占位]: '{p.text[:40]}'")

        # 检查2: 连续重复标题
        prev_heading_clean = None
        for p in doc.paragraphs:
            if p.style and p.style.name.startswith("BM_H"):
                clean = re.sub(r'\s+', '', strip_existing_number(p.text.strip()))
                if prev_heading_clean and clean == prev_heading_clean and clean:
                    error_issues.append(f"连续重复标题: '{p.text.strip()[:40]}'")
                prev_heading_clean = clean
            else:
                prev_heading_clean = None

        # 检查3: 标题后紧跟空段落
        heading_empty_count = 0
        for i in range(len(doc.paragraphs) - 1):
            curr = doc.paragraphs[i]
            nxt = doc.paragraphs[i + 1]
            if (curr.style and curr.style.name.startswith("BM_H")
                    and not nxt.text.strip()):
                heading_empty_count += 1
        if heading_empty_count > 0:
            error_issues.append(f"标题后空段落: {heading_empty_count} 处")

        # 检查4: [图片待人工确认] 残留统计
        pending_count = sum(
            1 for p in doc.paragraphs
            if "[图片待人工确认" in p.text or "[图片待确认" in p.text
        )
        if pending_count > 0:
            warning_issues.append(
                f"图片待人工确认: {pending_count} 处 -- "
                f"这些占位符需要人工审核后补充图片"
            )

        # 检查5: 图片记录数 vs 实际插入数差值
        if image_records:
            drawings = doc.element.body.findall(f'.//{qn("w:drawing")}')
            inserted = len(drawings)
            total = len(image_records)
            if inserted == 0:
                warning_issues.append(
                    f"图片未插入: 有 {total} 张图片记录但文档中无图片"
                    f" -- 请确认素材绑定状态"
                )
            elif inserted < total:
                gap = total - inserted
                warning_issues.append(
                    f"图片插入缺口: 记录={total}, 插入={inserted}, "
                    f"缺口={gap}"
                )

        if error_issues or warning_issues:
            for issue in error_issues:
                logger.error(f"渲染质检: {issue}")
            for issue in warning_issues:
                logger.warning(f"渲染质检: {issue}")
        else:
            logger.info("渲染质检: 全部通过")

    # ==================== 页眉页脚 ====================

    def _set_header_footer(
        self, doc: object, cfg: dict, Pt: type,
    ) -> None:
        """设置页眉页脚(含字体和页码)"""
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn

        header_text = sanitize_xml_compatible_text(cfg.get("header_text"))
        footer_text = sanitize_xml_compatible_text(cfg.get("footer_text"))

        for section in doc.sections:
            if header_text:
                header = section.header
                header.is_linked_to_previous = False
                if header.paragraphs:
                    hp = header.paragraphs[0]
                    hp.text = header_text
                    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in hp.runs:
                        run.font.size = Pt(HEADER_FOOTER_SIZE_PT)
                        rpr = run._r.find(qn("w:rPr"))
                        if rpr is not None:
                            _set_cjk_font(rpr, FONT_HEADER_FOOTER)

            if footer_text:
                footer = section.footer
                footer.is_linked_to_previous = False
                if footer.paragraphs:
                    fp = footer.paragraphs[0]
                    fp.text = footer_text
                    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in fp.runs:
                        run.font.size = Pt(HEADER_FOOTER_SIZE_PT)

    # ==================== 树操作 ====================

    def _flatten_tree(
        self, nodes: list[WritingOutlineNode],
    ) -> list[WritingOutlineNode]:
        """将树按阅读顺序展平"""
        node_map: dict[Optional[int], list[WritingOutlineNode]] = {}
        for n in nodes:
            node_map.setdefault(n.parent_id, []).append(n)

        flat: list[WritingOutlineNode] = []

        def _recurse(parent_id: Optional[int]) -> None:
            children = node_map.get(parent_id, [])
            children.sort(key=lambda n: n.sort_order)
            for child in children:
                flat.append(child)
                _recurse(child.id)

        _recurse(None)
        return flat

    # ==================== DB 读取 ====================

    async def _load_render_profile(
        self, project_id: int,
    ) -> Optional[WritingRenderProfile]:
        result = await self.db.execute(
            select(WritingRenderProfile).where(
                WritingRenderProfile.project_id == project_id,
            )
        )
        return result.scalars().first()

    async def _load_outline_tree(
        self, project_id: int,
    ) -> list[WritingOutlineNode]:
        result = await self.db.execute(
            select(WritingOutlineNode).where(
                WritingOutlineNode.project_id == project_id,
            ).order_by(WritingOutlineNode.depth, WritingOutlineNode.sort_order)
        )
        return list(result.scalars().all())

    async def _load_drafts_map(
        self, project_id: int,
    ) -> dict[int, WritingSectionDraft]:
        result = await self.db.execute(
            select(WritingSectionDraft).join(
                WritingOutlineNode,
                WritingSectionDraft.outline_node_id == WritingOutlineNode.id,
            ).where(
                WritingOutlineNode.project_id == project_id,
                WritingSectionDraft.is_current == 1,
            )
        )
        drafts = list(result.scalars().all())
        return {d.outline_node_id: d for d in drafts}

    async def _load_all_image_records(
        self, project_id: int,
    ) -> dict[int, WritingMaterialImage]:
        """预加载项目所有素材图片记录(按素材+序号稳定排序)"""
        from app.models.writing_material import WritingMaterial

        try:
            result = await self.db.execute(
                select(WritingMaterialImage).join(
                    WritingMaterial,
                    WritingMaterialImage.material_id == WritingMaterial.id,
                ).where(
                    WritingMaterial.project_id == project_id,
                ).order_by(
                    WritingMaterialImage.material_id,
                    WritingMaterialImage.image_index,
                    WritingMaterialImage.id,
                )
            )
            return {rec.id: rec for rec in result.scalars().all()}
        except Exception as e:
            logger.warning(f"预加载图片记录失败: {e}")
            return {}

    async def _build_node_image_map(
        self,
        project_id: int,
        image_records: dict[int, WritingMaterialImage],
    ) -> dict[int, dict[int, WritingMaterialImage]]:
        """构建章节级图片池，避免渲染阶段消费项目级全量图片。"""
        if not image_records:
            return {}

        result = await self.db.execute(
            select(WritingBinding).where(
                WritingBinding.project_id == project_id,
                WritingBinding.outline_node_id.isnot(None),
                WritingBinding.material_id.isnot(None),
                WritingBinding.binding_type == BindingType.material_to_node,
            )
        )
        bindings = list(result.scalars().all())
        if not bindings:
            return {}

        bindings_by_node: dict[int, list[WritingBinding]] = {}
        for binding in bindings:
            if binding.outline_node_id is None:
                continue
            bindings_by_node.setdefault(binding.outline_node_id, []).append(binding)

        material_to_images: dict[int, dict[int, WritingMaterialImage]] = {}
        for rec in image_records.values():
            material_to_images.setdefault(rec.material_id, {})[rec.id] = rec

        node_image_map: dict[int, dict[int, WritingMaterialImage]] = {}
        for node_id, node_bindings in bindings_by_node.items():
            # 渲染图片池与生成侧共享同一策略——素材绑定逐边只消费 confirmed。
            consumable = select_confirmed_material_bindings(node_bindings)
            if not consumable:
                continue
            node_images: dict[int, WritingMaterialImage] = {}
            for binding in consumable:
                if binding.material_id is None:
                    continue
                for rec_id, rec in material_to_images.get(binding.material_id, {}).items():
                    node_images[rec_id] = rec
            if node_images:
                node_image_map[node_id] = node_images
        return node_image_map


class HtmlPreviewBuilder:
    """HTML 预览生成器"""

    def __init__(self, db: AsyncSession) -> None:
        self.db = db

    async def build(self, project_id: int) -> str:
        """组装完整 HTML 预览"""
        logger.info(f"HTML预览生成启动: project={project_id}")
        renderer = DocRenderer(self.db)
        from app.core.writer.render_profile_currentness import ensure_render_profile_current
        await ensure_render_profile_current(self.db, project_id)
        profile = await renderer._load_render_profile(project_id)
        nodes = await renderer._load_outline_tree(project_id)
        drafts = await renderer._load_drafts_map(project_id)

        flat = renderer._flatten_tree(nodes)
        number_map = compute_chapter_numbers(nodes)
        cfg = merge_with_profile(profile)
        css = self._generate_css(cfg)

        # 图片为预览增强项: 加载失败时降级为无图预览, 不阻断正文渲染
        node_image_map: Optional[dict[int, dict[int, WritingMaterialImage]]] = None
        try:
            image_records = await renderer._load_all_image_records(project_id)
            node_image_map = await renderer._build_node_image_map(
                project_id, image_records,
            )
        except Exception as exc:
            logger.warning(
                f"预览图片加载失败,降级为无图预览: project={project_id}, err={exc}"
            )
            node_image_map = None

        # P4: 生成图示为预览增强项, 加载失败降级为占位标记
        diagram_records: dict[int, object] = {}
        try:
            from app.core.writer.visuals.repository import SectionDiagramRepository
            diagram_records = await SectionDiagramRepository(self.db).get_ready_by_project(
                project_id,
            )
        except Exception as exc:
            logger.warning(
                f"预览图示加载失败,降级为占位: project={project_id}, err={exc}"
            )

        body = self._render_body(flat, drafts, number_map, node_image_map, diagram_records)
        logger.info(f"HTML预览生成完成: project={project_id}, sections={len(flat)}")

        return (
            "<!DOCTYPE html>\n"
            "<html lang='zh-CN'>\n"
            f"<head><meta charset='utf-8'><style>{css}</style></head>\n"
            f"<body>{body}</body>\n"
            "</html>"
        )

    @staticmethod
    def _generate_css(cfg: dict) -> str:
        """根据合并后配置生成 CSS (与 Word 渲染对齐)"""
        fb = cfg.get("font_body", FONT_BODY)
        fh = cfg.get("font_heading", FONT_HEADING)
        sb = cfg.get("font_size_body", BODY_SIZE_PT)
        sh1 = cfg.get("font_size_h1", 16.0)
        sh2 = cfg.get("font_size_h2", 15.0)
        sh3 = cfg.get("font_size_h3", 14.0)
        ls = normalize_line_spacing(cfg.get("line_spacing"))
        mt = cfg.get("margin_top", MARGIN_TOP_MM)
        mb = cfg.get("margin_bottom", MARGIN_BOTTOM_MM)
        ml = cfg.get("margin_left", MARGIN_LEFT_MM)
        mr = cfg.get("margin_right", MARGIN_RIGHT_MM)

        h3_font = HEADING_SPECS["Heading 3"].font_name
        h4_font = HEADING_SPECS["Heading 4"].font_name
        h5_font = HEADING_SPECS["Heading 5"].font_name
        h5_ascii = HEADING_SPECS["Heading 5"].font_name_ascii
        h5_size = HEADING_SPECS["Heading 5"].size_pt

        return f"""
        body {{
            font-family: '{fb}', '宋体', SimSun, serif;
            font-size: {sb}pt;
            line-height: {ls};
            color: #000;
            max-width: 210mm;
            margin-left: auto;
            margin-right: auto;
            padding: {mt}mm {mr}mm {mb}mm {ml}mm;
        }}
        h1 {{ font-family: '{fh}', SimHei, sans-serif; font-size: {sh1}pt; margin-top: 24pt; color: #000; font-style: normal; }}
        h2 {{ font-family: '{fh}', SimHei, sans-serif; font-size: {sh2}pt; margin-top: 18pt; color: #000; font-style: normal; }}
        h3 {{ font-family: '{h3_font}', KaiTi, serif; font-size: {sh3}pt; margin-top: 14pt; color: #000; font-style: normal; }}
        h4 {{ font-family: '{h4_font}', FangSong, serif; font-size: {sb}pt; margin-top: 10pt; color: #000; font-style: normal; font-weight: bold; }}
        h5 {{ font-family: '{h5_font}', '{h5_ascii}', serif; font-size: {h5_size}pt; margin-top: 6pt; color: #000; font-style: normal; font-weight: bold; }}
        table {{ border-collapse: collapse; width: 100%; margin: 12pt 0; }}
        th, td {{ border: 1px solid #000; padding: 6pt 8pt; font-size: {sb - 1.5}pt; }}
        th {{ font-weight: bold; }}
        .section-divider {{ border-top: 1px dashed #ccc; margin: 24pt 0; }}
        .section-content > p {{ text-indent: 2em; margin: 0; }}
        .section-content .lead {{ text-indent: 0; font-weight: 700; margin: 4pt 0 0; }}
        table p {{ text-indent: 0; }}
        .placeholder {{ color: #c00; background: #fff3f3; padding: 2px 4px; }}
        .ai-mark {{ color: #e67e22; background: #fef9e7; padding: 2px 4px; border-radius: 2px; }}
        """

    def _render_body(
        self,
        flat_nodes: list[WritingOutlineNode],
        drafts: dict[int, WritingSectionDraft],
        number_map: Optional[dict[int, str]] = None,
        node_image_map: Optional[dict[int, dict[int, "WritingMaterialImage"]]] = None,
        diagram_records: Optional[dict[int, object]] = None,
    ) -> str:
        """渲染 HTML body"""
        parts: list[str] = []
        # 跨章节独占图片消费登记, 与 Word 渲染共享同一去重语义
        consumed_image_keys: dict[str, str] = {}

        for node in flat_nodes:
            tag = DocRenderer._get_html_tag_name(node.depth)
            number = (number_map or {}).get(node.id, "")
            clean_title = strip_existing_number(node.title)
            display_title = sanitize_xml_compatible_text(
                f"{number}{clean_title}" if number else clean_title,
            )
            parts.append(f"<{tag}>{display_title}</{tag}>")

            draft = drafts.get(node.id)
            if draft and draft.content:
                raw_html = draft.content.strip()
                content = DocRenderer._preprocess_html_for_bid(raw_html, clean_title)
                if is_fixed_template_node(node):
                    # 固定模板章节保真输出, 不走统一编号契约
                    rendered = render_verbatim_preview_html(self, content)
                else:
                    blocks = DocRenderer._merge_bare_numbering_with_following_paragraph(
                        parse_section_blocks(content),
                    )
                    rendered = self._render_section_blocks(blocks, node.depth)
                rendered = self._inline_node_images(
                    rendered, raw_html, node, node_image_map, consumed_image_keys,
                )
                # P4: 生成图示预览内联(与 Word 判定对齐)
                from app.core.writer.visuals.figure_bridge import (
                    inline_diagram_figures,
                )
                rendered = inline_diagram_figures(
                    rendered, raw_html, diagram_records or {},
                )
                parts.append(f'<div class="section-content">{rendered}</div>')
            else:
                parts.append(
                    f'<p class="placeholder">[{node.title} - 内容待补充]</p>'
                )

            parts.append('<div class="section-divider"></div>')

        return "\n".join(parts)

    @staticmethod
    def _inline_node_images(
        rendered_html: str,
        raw_html: str,
        node: WritingOutlineNode,
        node_image_map: Optional[dict[int, dict[int, "WritingMaterialImage"]]],
        consumed_image_keys: dict[str, str],
    ) -> str:
        """将章节渲染结果中的图片占位替换为 base64 内联图片。

        与 Word 渲染对齐: 先按 node.title 做亲和过滤得到候选图片池,
        再交由 preview_image_inliner 按占位顺序消费 (含 clean_title 闸门)。
        """
        if not node_image_map:
            return rendered_html
        node_images = node_image_map.get(node.id)
        if not node_images:
            return rendered_html

        from app.core.writer.image_node_affinity import filter_images_for_node
        from app.core.writer.preview_image_inliner import inline_section_images

        filtered = filter_images_for_node(list(node_images.values()), node.title)
        section_images = {img.id: img for img in filtered} if filtered else {}
        return inline_section_images(
            rendered_html,
            raw_html,
            section_images,
            strip_existing_number(node.title),
            consumed_image_keys,
        )

    def _render_section_blocks(
        self,
        blocks: list[SectionBlock],
        base_depth: int,
    ) -> str:
        parts: list[str] = []
        heading_counters: dict[int, int] = {}
        context_depth = base_depth

        for block in blocks:
            if block.kind == "heading":
                absolute_depth = DocRenderer._resolve_block_depth(
                    block,
                    preferred_depth=min(
                        base_depth + max(block.level, 1),
                        max(context_depth + 1, max(base_depth + 1, 1)),
                    ),
                    base_depth=base_depth,
                    source="preview_heading",
                )
                if block.pure_numbering:
                    prefix = DocRenderer._next_number_prefix(
                        absolute_depth,
                        heading_counters,
                    )
                    tag = DocRenderer._get_html_tag_name(absolute_depth)
                    parts.append(f"<{tag}>{escape(prefix)}</{tag}>")
                    context_depth = absolute_depth
                    continue
                clean_text = DocRenderer._sanitize_numbered_text(block.text, "preview_heading")
                if not clean_text:
                    context_depth = absolute_depth
                    continue
                prefix = DocRenderer._next_number_prefix(absolute_depth, heading_counters)
                tag = DocRenderer._get_html_tag_name(absolute_depth)
                parts.append(f"<{tag}>{escape(f'{prefix}{clean_text}')}</{tag}>")
                context_depth = absolute_depth
            elif block.kind == "numbered_paragraph":
                absolute_depth = DocRenderer._resolve_block_depth(
                    block,
                    preferred_depth=DocRenderer._compute_numbered_preferred_depth(
                        block,
                        base_depth,
                        context_depth,
                    ),
                    base_depth=base_depth,
                    source="preview_numbered_paragraph",
                )
                if block.pure_numbering:
                    prefix = DocRenderer._next_number_prefix(
                        absolute_depth,
                        heading_counters,
                    )
                    parts.append(f'<p class="lead">{escape(prefix)}</p>')
                    context_depth = absolute_depth
                    continue
                clean_text = DocRenderer._sanitize_numbered_text(
                    block.text, "preview_numbered_paragraph",
                )
                if not clean_text:
                    context_depth = absolute_depth
                    continue
                prefix = DocRenderer._next_number_prefix(absolute_depth, heading_counters)
                parts.append(f'<p class="lead">{escape(f"{prefix}{clean_text}")}</p>')
                context_depth = absolute_depth
            elif block.kind == "paragraph":
                parts.append(self._decorate_text_block(block.text))
            elif block.kind == "list":
                parts.append(self._render_list_html(block, context_depth))
            elif block.kind in {"table", "html"}:
                parts.append(block.html)

        return "\n".join(parts)

    def _render_list_html(
        self,
        block: SectionBlock,
        base_depth: int,
        nesting_level: int = 1,
    ) -> str:
        parts: list[str] = []
        for idx, item in enumerate(block.items):
            clean_text = DocRenderer._sanitize_numbered_text(item.text, "preview_list_item")
            if not clean_text:
                continue
            prefix = ""
            if block.ordered:
                prefix = format_number_for_depth(base_depth + nesting_level, idx)
            cls = "lead" if block.ordered else ""
            text = escape(f"{prefix}{clean_text}".strip())
            parts.append(f'<p class="{cls}">{text}</p>' if cls else f"<p>{text}</p>")
            for child in item.children:
                if child.kind == "list":
                    parts.append(self._render_list_html(child, base_depth, nesting_level + 1))
                elif child.kind == "paragraph":
                    parts.append(self._decorate_text_block(child.text))
                elif child.kind == "numbered_paragraph":
                    nested_text = DocRenderer._sanitize_numbered_text(
                        child.text, "preview_nested_numbered_paragraph",
                    )
                    if not nested_text:
                        continue
                    nested_base_depth = base_depth + nesting_level
                    nested_depth = DocRenderer._resolve_block_depth(
                        child,
                        preferred_depth=DocRenderer._compute_numbered_preferred_depth(
                            child,
                            nested_base_depth,
                            nested_base_depth,
                        ),
                        base_depth=nested_base_depth,
                        source="preview_nested_numbered_paragraph",
                    )
                    nested_prefix = format_number_for_depth(
                        nested_depth,
                        0,
                    )
                    parts.append(f'<p class="lead">{escape(f"{nested_prefix}{nested_text}")}</p>')
                elif child.kind == "heading":
                    child_depth = DocRenderer._resolve_block_depth(
                        child,
                        preferred_depth=min(
                            base_depth + nesting_level,
                            max(base_depth + nesting_level, 1),
                        ),
                        base_depth=base_depth + nesting_level - 1,
                        source="preview_nested_heading",
                    )
                    tag = DocRenderer._get_html_tag_name(child_depth)
                    clean_heading = DocRenderer._sanitize_numbered_text(
                        child.text, "preview_nested_heading",
                    )
                    if not clean_heading:
                        continue
                    parts.append(f"<{tag}>{escape(clean_heading)}</{tag}>")
                elif child.kind in {"table", "html"}:
                    parts.append(child.html)
        return "\n".join(parts)

    @staticmethod
    def _decorate_text_block(text: str) -> str:
        """段落文本装饰: 占位符替换 + 样式高亮: 旧实现仅对 {xxx} 加样式高亮但保留花括号字符,
 导致 Word 导出与 HTML 预览字符层不一致, 用户复制粘贴或导出后仍是 {xxx}。
 新实现先把 {xxx} 真正替换为 [待补充:xxx] (与 Word 路径对齐),
 再统一对 [待补充:xxx] 整体高亮, 保证字符层与样式层一致。
 """
        text = sanitize_residual_placeholders(text)
        content = escape(text)
        content = re.sub(
            r"\[待补充[^]]*\]",
            r'<span class="placeholder">\g<0></span>',
            content,
        )
        content = re.sub(
            r"\[!AI:[^\]]+\]",
            r'<span class="ai-mark">\g<0></span>',
            content,
        )
        return f"<p>{content}</p>"


def get_doc_renderer(db: AsyncSession) -> DocRenderer:
    """工厂函数"""
    return DocRenderer(db=db)


def get_html_preview_builder(db: AsyncSession) -> HtmlPreviewBuilder:
    """工厂函数"""
    return HtmlPreviewBuilder(db=db)

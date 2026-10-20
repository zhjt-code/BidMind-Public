"""身份证图片 Word 渲染助手。

职责: 将既定身份证布局计划写入 Word 无边框双列表格。
不做: 不推断图片语义、不查询数据库、不处理普通素材图片。
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional

from loguru import logger

from app.core.writer.identity_image_layout import (
    IdentityImageLayout,
    IdentityImageSlot,
)
from app.core.writer.image_node_affinity import is_exclusive_image
from app.core.xml_compat import sanitize_xml_compatible_text

_IDENTITY_CARD_WIDTH_MM = 78


def insert_identity_image_block(
    doc: object,
    anchor_paragraph: object,
    layout: IdentityImageLayout,
    captions_by_id: dict[int, str],
    Pt: type,
    consumed_image_keys: Optional[dict[str, str]],
    current_section_title: str,
) -> frozenset[int]:
    """在指定占位段落位置插入身份证双列表格，并返回已消费图片 ID。"""
    if not layout.rows:
        return frozenset()

    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _remove_table_borders(table)

    for identity_row in layout.rows:
        row = table.add_row()
        _keep_row_together(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _render_slot(
            row.cells[0], identity_row.front, captions_by_id, Pt,
            consumed_image_keys, current_section_title, WD_ALIGN_PARAGRAPH,
        )
        _render_slot(
            row.cells[1], identity_row.back, captions_by_id, Pt,
            consumed_image_keys, current_section_title, WD_ALIGN_PARAGRAPH,
        )
        for extra_slot in identity_row.unknown_slots:
            extra_row = table.add_row()
            _keep_row_together(extra_row)
            for cell in extra_row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _render_slot(
                extra_row.cells[0], extra_slot, captions_by_id, Pt,
                consumed_image_keys, current_section_title, WD_ALIGN_PARAGRAPH,
            )

    anchor = anchor_paragraph._element
    anchor.addnext(table._tbl)
    anchor_paragraph.clear()
    parent = anchor.getparent()
    if parent is not None:
        parent.remove(anchor)

    return layout.image_id_set


def _render_slot(
    cell: object,
    slot: IdentityImageSlot | None,
    captions_by_id: dict[int, str],
    Pt: type,
    consumed_image_keys: Optional[dict[str, str]],
    current_section_title: str,
    WD_ALIGN_PARAGRAPH: object,
) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if slot is None:
        return

    rec = slot.image
    caption = sanitize_xml_compatible_text(
        captions_by_id.get(slot.image_id) or slot.caption or "身份证图片",
    )

    if not Path(rec.image_path).exists():
        paragraph.add_run(f"[图片缺失: {caption}]").font.size = Pt(9)
        return

    if consumed_image_keys is not None and is_exclusive_image(rec):
        consume_key = rec.file_hash or rec.image_path
        prev_section = consumed_image_keys.get(consume_key)
        if prev_section and prev_section != current_section_title:
            paragraph.add_run(f"[图片待人工确认: {caption}]").font.size = Pt(9)
            return

    run = paragraph.add_run()
    try:
        from docx.shared import Mm

        run.add_picture(rec.image_path, width=Mm(_IDENTITY_CARD_WIDTH_MM))
    except Exception as exc:  # noqa: BLE001
        logger.warning("身份证图片插入失败: id={}, err={}", slot.image_id, exc)
        paragraph.add_run(f"[图片插入失败: {caption}]").font.size = Pt(9)
        return

    run.add_break()
    cap_run = paragraph.add_run(caption)
    cap_run.font.size = Pt(9)

    if consumed_image_keys is not None and is_exclusive_image(rec):
        consume_key = rec.file_hash or rec.image_path
        consumed_image_keys[consume_key] = current_section_title


def _remove_table_borders(table: object) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def _keep_row_together(row: object) -> None:
    from docx.oxml import OxmlElement

    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)

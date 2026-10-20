from tests.synthetic_data import synthetic_text
from types import SimpleNamespace
from pathlib import Path

from docx import Document
from docx.shared import Mm, Pt

from app.core.writer.bid_format_defaults import merge_with_profile
from app.core.writer.doc_renderer import DocRenderer


def test_apply_styles_creates_distinct_bm_heading_styles() -> None:
    doc = Document()
    renderer = DocRenderer(db=None)

    renderer._apply_styles(doc, merge_with_profile(None), Pt, Mm)

    h4 = doc.styles["BM_H4"]
    h5 = doc.styles["BM_H5"]

    assert h4.font.size != h5.font.size
    assert h4.paragraph_format.space_before != h5.paragraph_format.space_before
    assert h5.font.size.pt == 11.5


def _read_outline_level(style) -> str | None:
    from docx.oxml.ns import qn

    pPr = style.element.find(qn("w:pPr"))
    if pPr is None:
        return None
    outline = pPr.find(qn("w:outlineLvl"))
    if outline is None:
        return None
    return outline.get(qn("w:val"))


def test_apply_styles_sets_outline_levels_for_first_three_headings_only() -> None:
    """BM_H1/H2/H3 必须带大纲级别，供 Word「自动目录」收集；H4/H5 不进前三级。"""
    doc = Document()
    renderer = DocRenderer(db=None)

    renderer._apply_styles(doc, merge_with_profile(None), Pt, Mm)

    assert _read_outline_level(doc.styles["BM_H1"]) == "0"
    assert _read_outline_level(doc.styles["BM_H2"]) == "1"
    assert _read_outline_level(doc.styles["BM_H3"]) == "2"
    assert _read_outline_level(doc.styles["BM_H4"]) is None
    assert _read_outline_level(doc.styles["BM_H5"]) is None


def test_render_no_longer_inserts_toc_field_placeholder() -> None:
    """放弃半成品 TOC 域：样式具备大纲级别后，由用户在 Word 中自行生成目录。"""
    from app.core.writer import doc_renderer as module

    assert not hasattr(DocRenderer, "_insert_toc")
    source = Path(module.__file__).read_text(encoding="utf-8")
    assert "右键更新域" not in source
    assert "TOC \\h \\z \\u \\t" not in source


def test_html_to_docx_renders_explicit_numbering_with_restart() -> None:
    doc = Document()
    renderer = DocRenderer(db=None)
    renderer._apply_styles(doc, merge_with_profile(None), Pt, Mm)

    html = """
    <h4>服务目标</h4>
    <ol>
      <li>一级事项A
        <ol><li>子项A-1</li></ol>
      </li>
      <li>一级事项B
        <ol><li>子项B-1</li></ol>
      </li>
    </ol>
    """
    renderer._html_to_docx(
        doc,
        html,
        Pt,
        image_records={},
        section_title="测试章节",
        base_depth=1,
    )

    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "（一）服务目标" in texts
    assert "1、一级事项A" in texts
    assert "（1）子项A-1" in texts
    assert "2、一级事项B" in texts
    assert texts.count("（1）子项B-1") == 1


def test_preprocess_html_for_bid_removes_duplicate_plain_text_title() -> None:
    html = "项目方案\n<p>正文内容</p>"

    result = DocRenderer._preprocess_html_for_bid(html, "项目方案")

    assert "项目方案" not in result
    assert "正文内容" in result


def test_html_to_docx_strips_residual_double_number_text() -> None:
    doc = Document()
    renderer = DocRenderer(db=None)
    renderer._apply_styles(doc, merge_with_profile(None), Pt, Mm)

    html = """
    <ol>
      <li>1 投标函</li>
      <li>2 法定代表人身份证明及授权委托书</li>
    </ol>
    <p>3</p>
    """
    renderer._html_to_docx(
        doc,
        html,
        Pt,
        image_records={},
        section_title="测试章节",
        base_depth=1,
    )

    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "（一）1 投标函" not in texts
    assert "（二）2 法定代表人身份证明及授权委托书" not in texts
    assert "3、 2" not in texts
    assert "（一）投标函" in texts
    assert "（二）法定代表人身份证明及授权委托书" in texts


def test_html_to_docx_merges_pure_number_heading_with_following_paragraph() -> None:
    doc = Document()
    renderer = DocRenderer(db=None)
    renderer._apply_styles(doc, merge_with_profile(None), Pt, Mm)

    html = """
    <h4>（一）</h4>
    <p>我方具备完整的项目执行能力。</p>
    <h4>（二）</h4>
    <p>我方具备成熟的质量保障体系。</p>
    """
    renderer._html_to_docx(
        doc,
        html,
        Pt,
        image_records={},
        section_title="其他资料",
        base_depth=1,
    )

    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "（一）（一）" not in texts
    assert "（二）（二）" not in texts
    bare_one = [t for t in texts if t.strip() == "（一）"]
    bare_two = [t for t in texts if t.strip() == "（二）"]
    assert len(bare_one) == 0, "bare (一) should be merged with following paragraph"
    assert len(bare_two) == 0, "bare (二) should be merged with following paragraph"
    assert any("我方具备完整的项目执行能力" in t for t in texts)
    assert any("我方具备成熟的质量保障体系" in t for t in texts)


def test_html_to_docx_drops_template_symbol_and_punctuation_residue() -> None:
    doc = Document()
    renderer = DocRenderer(db=None)
    renderer._apply_styles(doc, merge_with_profile(None), Pt, Mm)

    html = """
    <p>¥</p>
    <p>） 我方将按招标文件要求提供资料。</p>
    <p>（一）.</p>
    <p>（三）. 60</p>
    """
    renderer._html_to_docx(
        doc,
        html,
        Pt,
        image_records={},
        section_title="测试章节",
        base_depth=1,
    )

    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "¥" not in texts
    assert "） 我方将按招标文件要求提供资料。" not in texts
    assert any("我方将按招标文件要求提供资料。" in text for text in texts)
    assert "（一）." not in texts
    assert "（三）. 60" not in texts


def test_html_to_docx_prefers_context_depth_when_visible_numbering_overshoots() -> None:
    doc = Document()
    renderer = DocRenderer(db=None)
    renderer._apply_styles(doc, merge_with_profile(None), Pt, Mm)

    html = "<h4>（1）核心资质与专业认证</h4>"
    renderer._html_to_docx(
        doc,
        html,
        Pt,
        image_records={},
        section_title="公司基本情况介绍",
        base_depth=2,
    )

    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "1、核心资质与专业认证" in texts
    assert "（1）核心资质与专业认证" not in texts


def test_html_to_docx_collapses_first_h5_heading_that_skips_a_level() -> None:
    doc = Document()
    renderer = DocRenderer(db=None)
    renderer._apply_styles(doc, merge_with_profile(None), Pt, Mm)

    html = "<h5>（1）核心资质与专业认证</h5>"
    renderer._html_to_docx(
        doc,
        html,
        Pt,
        image_records={},
        section_title="公司基本情况介绍",
        base_depth=2,
    )

    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "1、核心资质与专业认证" in texts
    assert "（1）核心资质与专业认证" not in texts


def test_html_to_docx_keeps_arabic_children_under_parenthesized_parent() -> None:
    doc = Document()
    renderer = DocRenderer(db=None)
    renderer._apply_styles(doc, merge_with_profile(None), Pt, Mm)

    html = """
    <h4>（九）承诺及声明</h4>
    <p>1、我方将按招标文件要求履行投标责任。</p>
    <p>2、本投标自开标日起有效期为60个日历日。</p>
    """
    renderer._html_to_docx(
        doc,
        html,
        Pt,
        image_records={},
        section_title="投标函",
        base_depth=1,
    )

    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "（一）承诺及声明" in texts
    assert "1、我方将按招标文件要求履行投标责任。" in texts
    assert "2、本投标自开标日起有效期为60个日历日。" in texts
    assert "（十）我方将按招标文件要求履行投标责任。" not in texts


def test_html_to_docx_removes_duplicate_plain_text_fixed_template_title() -> None:
    doc = Document()
    renderer = DocRenderer(db=None)
    renderer._apply_styles(doc, merge_with_profile(None), Pt, Mm)

    html = "（一）投标函\n（一）法定代表人身份证明及授权委托书"
    renderer._html_to_docx(
        doc,
        html,
        Pt,
        image_records={},
        section_title="投标函",
        base_depth=1,
    )

    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    assert "（一）投标函" not in texts
    assert "（一）法定代表人身份证明及授权委托书" in texts


def test_html_to_docx_merges_bare_numbering_with_following_paragraph() -> None:
    doc = Document()
    renderer = DocRenderer(db=None)
    renderer._apply_styles(doc, merge_with_profile(None), Pt, Mm)

    html = """
    <h4>（三）</h4>
    <p>我公司自成立以来，始终专注于音视频系统集成领域。</p>
    <h4>（四）项目团队</h4>
    <p>为确保本项目高标准高质量完成。</p>
    """
    renderer._html_to_docx(
        doc,
        html,
        Pt,
        image_records={},
        section_title="投标人提交的其他资料",
        base_depth=1,
    )

    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    bare_three = [t for t in texts if t.strip() == "（三）"]
    assert len(bare_three) == 0, f"bare (三) should be merged, got {texts}"
    assert any("我公司自成立以来" in t and "（" in t for t in texts)
    assert any("项目团队" in t for t in texts)


def test_verify_render_output_demotes_pending_image_and_gap_to_warning(monkeypatch) -> None:
    doc = Document()
    doc.add_paragraph("第一章 测试章节", style="Heading 1")
    doc.add_paragraph("[图片待确认: 未匹配到章节相关图片, 章节=测试章节]")

    warnings: list[str] = []
    errors: list[str] = []

    monkeypatch.setattr(
        "app.core.writer.doc_renderer.logger.warning",
        lambda message: warnings.append(message),
    )
    monkeypatch.setattr(
        "app.core.writer.doc_renderer.logger.error",
        lambda message: errors.append(message),
    )

    DocRenderer._verify_render_output(
        doc,
        {
            1: SimpleNamespace(id=1),
            2: SimpleNamespace(id=2),
        },
    )

    assert any(
        ("图片待确认" in msg) or ("图片待人工确认" in msg)
        for msg in warnings
    )
    assert any(
        ("图片未插入" in msg) or ("图片插入缺口" in msg)
        for msg in warnings
    )
    assert not any("图片待确认" in msg for msg in errors)
    assert not any("图片插入缺口" in msg for msg in errors)


def test_insert_figure_images_marks_pending_when_image_not_in_section_pool() -> None:
    doc = Document()
    paragraph = doc.add_paragraph("[图片占位]")

    DocRenderer._insert_figure_images(
        doc,
        '<figure data-image-id="1"><figcaption>图1: 税务登记证</figcaption></figure>',
        Pt,
        image_records={},
        start_para_index=0,
        current_section_title="税务登记证",
        consumed_image_keys={},
    )

    assert "图片待人工确认" in paragraph.text


def test_insert_figure_images_blocks_duplicate_exclusive_image_across_sections(monkeypatch) -> None:
    from docx.text.run import Run
    from pathlib import Path

    monkeypatch.setattr(Path, "exists", lambda self: True)
    monkeypatch.setattr(Run, "add_picture", lambda self, path, width=None: None)

    doc = Document()
    first_para = doc.add_paragraph("[图片占位]")
    image_record = SimpleNamespace(
        id=1,
        image_path="fake.jpg",
        file_hash="same-image",
        image_type=None,
        business_subtype="id_card",
        caption="图1: 法定代表人身份证",
        width=800,
        height=600,
    )
    consumed: dict[str, str] = {}

    DocRenderer._insert_figure_images(
        doc,
        '<figure data-image-id="1"><figcaption>图1: 法定代表人身份证</figcaption></figure>',
        Pt,
        image_records={1: image_record},
        start_para_index=0,
        current_section_title="法定代表人身份证明",
        consumed_image_keys=consumed,
    )

    second_para = doc.add_paragraph("[图片占位]")
    DocRenderer._insert_figure_images(
        doc,
        '<figure data-image-id="1"><figcaption>图1: 法定代表人身份证</figcaption></figure>',
        Pt,
        image_records={1: image_record},
        start_para_index=len(doc.paragraphs) - 1,
        current_section_title="HSE与质量保障承诺",
        consumed_image_keys=consumed,
    )

    assert first_para.text == ""
    assert "图片待人工确认" in second_para.text


def test_insert_figure_images_centers_image_and_clears_body_indent(monkeypatch) -> None:
    """: 素材图段落须居中并清零首行缩进，避免 BM_BODY 缩进导致右偏。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt as DocxPt
    from docx.text.run import Run
    from pathlib import Path

    monkeypatch.setattr(Path, "exists", lambda self: True)
    monkeypatch.setattr(Run, "add_picture", lambda self, path, width=None: None)

    doc = Document()
    DocRenderer(db=None)._apply_styles(doc, merge_with_profile(None), Pt, Mm)
    paragraph = doc.add_paragraph("[图片占位]", style="BM_BODY")
    image_record = SimpleNamespace(
        id=42,
        image_path="cert.jpg",
        file_hash="cert-hash",
        image_type=None,
        business_subtype=None,
        caption="图1: 营业执照",
        width=1200,
        height=900,
    )

    DocRenderer._insert_figure_images(
        doc,
        '<figure data-image-id="42"><figcaption>图1: 营业执照</figcaption></figure>',
        Pt,
        image_records={42: image_record},
        start_para_index=0,
        current_section_title="营业执照",
        consumed_image_keys={},
    )

    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paragraph.paragraph_format.first_line_indent == DocxPt(0)
    caption_para = None
    for candidate in doc.paragraphs:
        if candidate.text == "图1: 营业执照":
            caption_para = candidate
            break
    assert caption_para is not None
    assert caption_para.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert caption_para.paragraph_format.first_line_indent == DocxPt(0)


def test_insert_figure_images_allows_later_placeholder_after_first_image_missing(monkeypatch) -> None:
    from docx.text.run import Run
    from pathlib import Path

    monkeypatch.setattr(Path, "exists", lambda self: self.name != "missing.jpg")
    monkeypatch.setattr(Run, "add_picture", lambda self, path, width=None: None)

    doc = Document()
    first_para = doc.add_paragraph("[图片占位]")
    second_para = doc.add_paragraph("[图片占位]")
    image_records = {
        1: SimpleNamespace(
            id=1,
            image_path="missing.jpg",
            file_hash="missing",
            image_type=None,
            business_subtype="tax_registration",
            caption="图1: 税务登记证",
            width=800,
            height=600,
        ),
        2: SimpleNamespace(
            id=2,
            image_path="ok.jpg",
            file_hash="ok",
            image_type=None,
            business_subtype="tax_registration",
            caption="图2: 税务登记证副本",
            width=800,
            height=600,
        ),
    }

    DocRenderer._insert_figure_images(
        doc,
        (
            '<figure data-image-id="1"><figcaption>图1: 税务登记证</figcaption></figure>'
            '<figure data-image-id="2"><figcaption>图2: 税务登记证副本</figcaption></figure>'
        ),
        Pt,
        image_records=image_records,
        start_para_index=0,
        current_section_title="税务登记证",
        consumed_image_keys={},
    )

    assert "图片缺失" in first_para.text
    assert second_para.text == ""


def test_insert_figure_images_renders_identity_cards_as_two_column_block(monkeypatch) -> None:
    from docx.text.run import Run
    from pathlib import Path

    from app.models.writing_material_image import ImageType

    monkeypatch.setattr(Path, "exists", lambda self: True)
    inserted_paths: list[str] = []
    monkeypatch.setattr(
        Run,
        "add_picture",
        lambda self, path, width=None: inserted_paths.append(path),
    )

    doc = Document()
    for _ in range(4):
        doc.add_paragraph("[图片占位]")

    image_records = {
        4: SimpleNamespace(
            id=4,
            image_path="agent-back.jpg",
            file_hash="agent-back",
            image_type=ImageType.personnel_cert,
            business_subtype="id_card",
            caption="委托代理人身份证反面",
            vl_content="签发机关 XX公安局 有效期限 2020-2030",
            structured_fields="{}",
            material_id=1,
            image_index=3,
            width=1200,
            height=760,
        ),
        2: SimpleNamespace(
            id=2,
            image_path="legal-back.jpg",
            file_hash="legal-back",
            image_type=ImageType.personnel_cert,
            business_subtype="id_card",
            caption="法定代表人身份证反面",
            vl_content="签发机关 XX公安局 有效期限 2020-2030",
            structured_fields="{}",
            material_id=1,
            image_index=1,
            width=1200,
            height=760,
        ),
        3: SimpleNamespace(
            id=3,
            image_path="agent-front.jpg",
            file_hash="agent-front",
            image_type=ImageType.personnel_cert,
            business_subtype="id_card",
            caption="委托代理人身份证正面",
            vl_content="",
            structured_fields=synthetic_text('{"姓名": "<PERSON_B>", "证件号码": "', '110101', '199001', '011234', '"}'),
            material_id=1,
            image_index=2,
            width=1200,
            height=760,
        ),
        1: SimpleNamespace(
            id=1,
            image_path="legal-front.jpg",
            file_hash="legal-front",
            image_type=ImageType.personnel_cert,
            business_subtype="id_card",
            caption="法定代表人身份证正面",
            vl_content="",
            structured_fields=synthetic_text('{"姓名": "<PERSON_A>", "证件号码": "', '110101', '198001', '011234', '"}'),
            material_id=1,
            image_index=0,
            width=1200,
            height=760,
        ),
    }

    DocRenderer._insert_figure_images(
        doc,
        (
            '<figure data-image-id="4"><figcaption>委托代理人身份证反面</figcaption></figure>'
            '<figure data-image-id="2"><figcaption>法定代表人身份证反面</figcaption></figure>'
            '<figure data-image-id="3"><figcaption>委托代理人身份证正面</figcaption></figure>'
            '<figure data-image-id="1"><figcaption>法定代表人身份证正面</figcaption></figure>'
        ),
        Pt,
        image_records=image_records,
        start_para_index=0,
        current_section_title="法定代表人身份证明及授权委托书",
        consumed_image_keys={},
    )

    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 2
    assert inserted_paths == [
        "legal-front.jpg",
        "legal-back.jpg",
        "agent-front.jpg",
        "agent-back.jpg",
    ]
    assert not any("[图片占位]" in paragraph.text for paragraph in doc.paragraphs)


# ==================== line_spacing 归一化回归测试 ====================


def test_normalize_line_spacing_preserves_valid_range() -> None:
    from app.core.writer.bid_format_defaults import normalize_line_spacing, LINE_SPACING
    assert normalize_line_spacing(1.0) == 1.0
    assert normalize_line_spacing(1.5) == 1.5
    assert normalize_line_spacing(2.0) == 2.0
    assert normalize_line_spacing(3.0) == 3.0


def test_normalize_line_spacing_rejects_out_of_range() -> None:
    from app.core.writer.bid_format_defaults import normalize_line_spacing, LINE_SPACING
    assert normalize_line_spacing(22) == LINE_SPACING
    assert normalize_line_spacing(22.0) == LINE_SPACING
    assert normalize_line_spacing(0) == LINE_SPACING
    assert normalize_line_spacing(0.5) == LINE_SPACING
    assert normalize_line_spacing(-1) == LINE_SPACING
    assert normalize_line_spacing(100) == LINE_SPACING
    assert normalize_line_spacing(3.1) == LINE_SPACING


def test_normalize_line_spacing_handles_none_and_garbage() -> None:
    from app.core.writer.bid_format_defaults import normalize_line_spacing, LINE_SPACING
    assert normalize_line_spacing(None) == LINE_SPACING
    assert normalize_line_spacing("abc") == LINE_SPACING
    assert normalize_line_spacing("") == LINE_SPACING
    assert normalize_line_spacing([]) == LINE_SPACING


def test_apply_styles_normalizes_line_spacing_22() -> None:
    """ 复现: line_spacing=22 不得生成 ~22 倍行距"""
    doc = Document()
    renderer = DocRenderer(db=None)

    bad_cfg = merge_with_profile(None)
    bad_cfg["line_spacing"] = 22.0

    renderer._apply_styles(doc, bad_cfg, Pt, Mm)

    normal = doc.styles["Normal"]
    assert normal.paragraph_format.line_spacing == 1.5, (
        f"Expected 1.5, got {normal.paragraph_format.line_spacing}"
    )
    body = doc.styles["BM_BODY"]
    assert body.paragraph_format.line_spacing == 1.5
    h1 = doc.styles["BM_H1"]
    assert h1.paragraph_format.line_spacing == 1.5


def test_merge_with_profile_normalizes_line_spacing() -> None:
    """merge_with_profile 出口也应归一化"""
    profile = SimpleNamespace(
        font_body=None, font_heading=None,
        font_size_body=None, font_size_h1=None, font_size_h2=None, font_size_h3=None,
        line_spacing=22.0,
        margin_top=None, margin_bottom=None, margin_left=None, margin_right=None,
        header_text=None, footer_text=None,
    )
    cfg = merge_with_profile(profile)
    assert cfg["line_spacing"] == 1.5


# ====================: 复杂表格兼容性测试 ====================

# htmldocx 按首行列数创建 python-docx Table；后续行列数超出即 IndexError。
# 多组 thead/tbody 分段中后段列数 > 前段列数时必然触发。
_BUG165_MULTI_GROUP_HTML = '<table>\n<thead><tr><th>拟投入人数</th><th>要求人数</th></tr></thead>\n<tbody><tr><td>15人</td><td>不少于10人</td></tr></tbody>\n<thead><tr><th>序号</th><th>姓名</th><th>学历</th><th>专业</th><th>职称</th><th>岗位证书</th></tr></thead>\n<tbody>\n<tr><td>1</td><td><PERSON_A></td><td>本科</td><td>土木工程</td><td>高级工程师</td><td>一级建造师</td></tr>\n</tbody>\n</table>'

_BUG165_SINGLE_EXCESS_HTML = """\
<table>
<thead><tr><th>Col1</th></tr></thead>
<tbody><tr><td>A</td><td>B</td></tr></tbody>
</table>"""

_STANDARD_TABLE_HTML = """\
<table>
<thead><tr><th>Name</th><th>Role</th></tr></thead>
<tbody><tr><td>Alice</td><td>PM</td></tr><tr><td>Bob</td><td>Dev</td></tr></tbody>
</table>"""


def test_html_to_docx_handles_multi_thead_tbody_as_real_table(monkeypatch) -> None:
    """: 多组 thead/tbody 且后段列数多于前段应渲染为表格而非降级纯文本。"""
    doc = Document()
    renderer = DocRenderer(db=None)
    renderer._apply_styles(doc, merge_with_profile(None), Pt, Mm)

    errors: list[str] = []
    monkeypatch.setattr(
        "app.core.writer.doc_renderer.logger.error",
        lambda message: errors.append(message),
    )

    renderer._html_to_docx(
        doc,
        _BUG165_MULTI_GROUP_HTML,
        Pt,
        image_records={},
        section_title="团队主要人员名单与数量证明",
        base_depth=2,
    )

    assert len(doc.tables) >= 1, "should produce at least one real table"
    all_cells = [
        cell.text for tbl in doc.tables for row in tbl.rows for cell in row.cells
    ]
    assert any('<PERSON_A>' in c for c in all_cells), "table cell content preserved"
    assert not any("降级为纯文本" in e for e in errors), (
        "should NOT trigger fallback-to-plain-text error log"
    )


def test_html_to_docx_handles_row_exceeding_thead_cols(monkeypatch) -> None:
    """tbody 行列数超过 thead 列数应渲染为表格而非降级纯文本。"""
    doc = Document()
    renderer = DocRenderer(db=None)
    renderer._apply_styles(doc, merge_with_profile(None), Pt, Mm)

    errors: list[str] = []
    monkeypatch.setattr(
        "app.core.writer.doc_renderer.logger.error",
        lambda message: errors.append(message),
    )

    renderer._html_to_docx(
        doc,
        _BUG165_SINGLE_EXCESS_HTML,
        Pt,
        image_records={},
        section_title="test",
        base_depth=1,
    )

    assert len(doc.tables) >= 1, "should produce a real table"
    assert not any("降级为纯文本" in e for e in errors)


def test_html_to_docx_standard_table_still_works() -> None:
    """标准表格回归: 正常结构不受归一化影响。"""
    doc = Document()
    renderer = DocRenderer(db=None)
    renderer._apply_styles(doc, merge_with_profile(None), Pt, Mm)

    renderer._html_to_docx(
        doc,
        _STANDARD_TABLE_HTML,
        Pt,
        image_records={},
        section_title="test",
        base_depth=1,
    )

    assert len(doc.tables) >= 1
    cells = [cell.text for row in doc.tables[0].rows for cell in row.cells]
    assert "Alice" in cells
    assert "Bob" in cells


# ==================== 归一化函数单元测试 ====================


def test_normalize_table_html_splits_multi_thead_tbody() -> None:
    """多组 thead/tbody 拆成多张 table。"""
    from app.core.writer.doc_renderer import _normalize_table_html_for_docx
    from bs4 import BeautifulSoup

    html = (
        "<table>"
        "<thead><tr><th>A</th></tr></thead>"
        "<tbody><tr><td>1</td></tr></tbody>"
        "<thead><tr><th>B</th><th>C</th></tr></thead>"
        "<tbody><tr><td>2</td><td>3</td></tr></tbody>"
        "</table>"
    )
    result = _normalize_table_html_for_docx(html)
    soup = BeautifulSoup(result, "html.parser")
    tables = soup.find_all("table")
    assert len(tables) == 2, f"expected 2 tables, got {len(tables)}"


def test_normalize_table_html_pads_columns() -> None:
    """tbody 列数超过 thead 时，thead 行被补齐。"""
    from app.core.writer.doc_renderer import _normalize_table_html_for_docx
    from bs4 import BeautifulSoup

    html = (
        "<table>"
        "<thead><tr><th>A</th></tr></thead>"
        "<tbody><tr><td>1</td><td>2</td></tr></tbody>"
        "</table>"
    )
    result = _normalize_table_html_for_docx(html)
    soup = BeautifulSoup(result, "html.parser")
    thead_cells = soup.find("thead").find("tr").find_all(["th", "td"])
    assert len(thead_cells) == 2, f"expected 2 header cells, got {len(thead_cells)}"


def test_normalize_table_html_leaves_standard_table_unchanged() -> None:
    """标准表格不做任何修改。"""
    from app.core.writer.doc_renderer import _normalize_table_html_for_docx

    html = (
        "<table>"
        "<thead><tr><th>A</th><th>B</th></tr></thead>"
        "<tbody><tr><td>1</td><td>2</td></tr></tbody>"
        "</table>"
    )
    result = _normalize_table_html_for_docx(html)
    assert result == html, "standard table should pass through unchanged"


def test_normalize_table_html_handles_non_table_input() -> None:
    """非表格 HTML 原样返回。"""
    from app.core.writer.doc_renderer import _normalize_table_html_for_docx

    html = "<p>hello</p>"
    assert _normalize_table_html_for_docx(html) == html


def test_html_to_docx_table_inside_list_does_not_crash(monkeypatch) -> None:
    """列表子块内的复杂表格不应导致渲染崩溃。"""
    doc = Document()
    renderer = DocRenderer(db=None)
    renderer._apply_styles(doc, merge_with_profile(None), Pt, Mm)

    errors: list[str] = []
    monkeypatch.setattr(
        "app.core.writer.doc_renderer.logger.error",
        lambda message: errors.append(message),
    )

    html = """
    <ul>
      <li>项目概况
        <table>
        <thead><tr><th>Col1</th></tr></thead>
        <tbody><tr><td>A</td><td>B</td></tr></tbody>
        </table>
      </li>
    </ul>
    """
    renderer._html_to_docx(
        doc,
        html,
        Pt,
        image_records={},
        section_title="test",
        base_depth=1,
    )

    assert not any("降级为纯文本" in e for e in errors), (
        "list-nested table should not trigger fallback"
    )

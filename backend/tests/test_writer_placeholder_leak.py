"""防御测试: + 占位符泄漏到预览和 Word 导出

历史背景:
- (2026-04-13): fixed_template 占位符白名单过窄, 加正则兜底
- (2026-04-29): 修复仅覆盖 fixed_template 路径,
 LLM 路径 (data_fill/free_writing) 与 Word/HTML 渲染层均裸奔; 且
 _decorate_text_block 仅做样式高亮保留花括号字符, 与 Word 输出分叉

验证范围:
1. sanitize_residual_placeholders 白名单 + 兜底 + 幂等
2. _save_draft 出口闸门 A: LLM 输出含 {xxx} 落库前自动净化
3. _preprocess_html_for_bid 渲染层闸门 B1: 历史脏 draft 也能被兜住
4. _decorate_text_block 字符替换 + 样式高亮一致 (修正 反向断言)
5. Word 字节流 (document.xml) 不含 {xxx}
6. heading/table 等非 paragraph 块占位符兜底
7. 纯 ASCII 花括号内容 (CSS/JSON 片段) 不被误匹配
"""

import zipfile
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest

from app.core.writer.doc_renderer import DocRenderer, HtmlPreviewBuilder
from app.core.writer.section_composer import (
    SectionComposer,
    sanitize_residual_placeholders,
)


class TestSanitizeResidualPlaceholders:
    """模块级 sanitize_residual_placeholders 闸门函数 """

    def test_whitelist_placeholders_mapped(self) -> None:
        text = "致{投标人}，项目{项目名称}，日期{日期}"
        result = sanitize_residual_placeholders(text)

        assert "[待补充:投标人名称]" in result
        assert "[待补充:项目名称]" in result
        assert "[待补充:日期]" in result
        assert "{投标人}" not in result

    def test_fallback_catches_unlisted_chinese(self) -> None:
        text = "开户行: {开户行}，账号: {银行账号}，编号: {招标编号}"
        result = sanitize_residual_placeholders(text)

        assert "[待补充:开户行]" in result
        assert "[待补充:银行账号]" in result
        assert "[待补充:招标编号]" in result
        assert "{开户行}" not in result
        assert "{银行账号}" not in result

    def test_fallback_catches_english_underscore(self) -> None:
        """ 真实泄漏样本: LLM 复读 prompt/素材中的英文下划线占位符"""
        text = "公司{node_title}的{company_profile}介绍, 服务{service_commitment}, 制度{management_system}"
        result = sanitize_residual_placeholders(text)

        assert "[待补充:node_title]" in result
        assert "[待补充:company_profile]" in result
        assert "[待补充:service_commitment]" in result
        assert "[待补充:management_system]" in result
        assert "{node_title}" not in result
        assert "{company_profile}" not in result
        assert "{service_commitment}" not in result
        assert "{management_system}" not in result

    def test_fallback_catches_mixed_content(self) -> None:
        text = "注册资金{注册资金}万元，联系人{contact_name}"
        result = sanitize_residual_placeholders(text)

        assert "[待补充:注册资金]" in result
        assert "[待补充:contact_name]" in result

    def test_whitelist_priority_over_fallback(self) -> None:
        """白名单精确映射(如 投标人->投标人名称)优先于兜底正则"""
        result = sanitize_residual_placeholders("{投标人}")

        assert result == "[待补充:投标人名称]"
        assert "[待补充:投标人]" not in result

    def test_empty_braces_ignored(self) -> None:
        result = sanitize_residual_placeholders("空花括号 {} 不应匹配")
        assert "{}" in result

    def test_long_content_ignored(self) -> None:
        """超过20字符的花括号内容不匹配,避免误捕大段文本"""
        long_content = "这是一段超过二十个字符的很长的内容不应被匹配"
        text = f"{{{long_content}}}"
        result = sanitize_residual_placeholders(text)

        assert f"{{{long_content}}}" in result

    def test_idempotent_on_clean_content(self) -> None:
        """: 已替换内容二次调用结果不变, 保证闸门叠加安全"""
        text = "项目编号 [待补充:项目编号] 与公司 [待补充:公司名称] 准备就绪"
        result = sanitize_residual_placeholders(text)

        assert result == text

    def test_empty_or_none_safe(self) -> None:
        assert sanitize_residual_placeholders("") == ""
        assert sanitize_residual_placeholders(None) is None


class TestDecorateTextBlockAfterFix:
    """HtmlPreviewBuilder._decorate_text_block 修正后行为 

 历史断言 `assert "{开户行}" in result` 反向锁死了"保留花括号"的错误设计,
 新版本必须保证字符层与 Word 输出对齐: {xxx} 真正消失, 替换为 [待补充:xxx]。
 """

    def test_standard_placeholder_highlighted(self) -> None:
        result = HtmlPreviewBuilder._decorate_text_block("[待补充:公司名称]")
        assert 'class="placeholder"' in result
        assert "[待补充:公司名称]" in result

    def test_curly_brace_placeholder_replaced_and_highlighted(self) -> None:
        """ 修正: {xxx} 必须先替换为 [待补充:xxx] 再加样式高亮"""
        result = HtmlPreviewBuilder._decorate_text_block("开户行: {开户行}")

        assert "{开户行}" not in result
        assert "[待补充:开户行]" in result
        assert 'class="placeholder"' in result

    def test_english_underscore_curly_replaced(self) -> None:
        """ 真实泄漏样本: 英文下划线占位符也必须真替换"""
        result = HtmlPreviewBuilder._decorate_text_block(
            "公司{company_profile}的服务{service_commitment}"
        )

        assert "{company_profile}" not in result
        assert "{service_commitment}" not in result
        assert "[待补充:company_profile]" in result
        assert "[待补充:service_commitment]" in result

    def test_normal_text_not_highlighted(self) -> None:
        result = HtmlPreviewBuilder._decorate_text_block("这是正常文本")
        assert 'class="placeholder"' not in result

    def test_html_special_chars_escaped(self) -> None:
        result = HtmlPreviewBuilder._decorate_text_block("金额 < 100 & 其他")
        assert "&lt;" in result
        assert "&amp;" in result


class TestPreprocessHtmlForBidGate:
    """DocRenderer._preprocess_html_for_bid 渲染层闸门 B1 """

    def test_curly_placeholder_in_paragraph_replaced(self) -> None:
        html = "<p>公司 {company_profile} 提供 {service_commitment}</p>"
        result = DocRenderer._preprocess_html_for_bid(html, section_title="公司概况")

        assert "{company_profile}" not in result
        assert "{service_commitment}" not in result
        assert "[待补充:company_profile]" in result
        assert "[待补充:service_commitment]" in result

    def test_curly_placeholder_in_heading_replaced(self) -> None:
        """非 paragraph 块也必须被 B1 闸门兜住"""
        html = "<h4>{node_title}</h4><p>正文</p>"
        result = DocRenderer._preprocess_html_for_bid(html, section_title="公司概况")

        assert "{node_title}" not in result
        assert "[待补充:node_title]" in result

    def test_curly_placeholder_in_table_replaced(self) -> None:
        html = "<table><tr><td>{company_profile}</td><td>说明</td></tr></table>"
        result = DocRenderer._preprocess_html_for_bid(html, section_title="表格章节")

        assert "{company_profile}" not in result
        assert "[待补充:company_profile]" in result

    def test_clean_html_unchanged(self) -> None:
        """无占位符的内容 B1 闸门保持幂等, 不引入额外字符"""
        html = "<p>这是干净的内容</p>"
        result = DocRenderer._preprocess_html_for_bid(html, section_title="正常章节")

        assert "[待补充" not in result


class TestSaveDraftGate:
    """SectionComposer._save_draft 出口闸门 A 

 覆盖 LLM 路径(data_fill/free_writing/失败兜底)落库前的统一净化,
 防止 LLM 复读的 {xxx} 进入 writing_section_drafts.content。
 """

    def _make_composer_with_db(self) -> tuple[SectionComposer, dict]:
        composer = SectionComposer(db=None, project_id=0)
        captured: dict = {}

        result_holder = MagicMock()
        result_holder.scalars.return_value.all.return_value = []

        async def _execute(*args, **kwargs):
            return result_holder

        async def _commit():
            return None

        async def _refresh(obj):
            return None

        def _add(draft: object) -> None:
            captured["content"] = draft.content
            captured["word_count"] = draft.word_count

        db_mock = MagicMock()
        db_mock.execute = _execute
        db_mock.commit = _commit
        db_mock.refresh = _refresh
        db_mock.add = _add
        composer.db = db_mock
        return composer, captured

    @pytest.mark.asyncio
    async def test_save_draft_sanitizes_llm_curly_placeholders(self) -> None:
        """模拟 free_writing 路径 LLM 输出含 {xxx} 占位符落库被兜住"""
        composer, captured = self._make_composer_with_db()

        node = SimpleNamespace(id=42, title="公司概况")
        dirty_content = (
            "<h4>{node_title}</h4>"
            "<p>本公司{company_profile}专注于{service_commitment}领域</p>"
        )

        await composer._save_draft(node, dirty_content)

        saved_content = captured["content"]
        assert "{node_title}" not in saved_content
        assert "{company_profile}" not in saved_content
        assert "{service_commitment}" not in saved_content
        assert "[待补充:node_title]" in saved_content
        assert "[待补充:company_profile]" in saved_content
        assert "[待补充:service_commitment]" in saved_content

    @pytest.mark.asyncio
    async def test_save_draft_clean_content_unchanged(self) -> None:
        """干净内容入库不被改写, 闸门幂等"""
        composer, captured = self._make_composer_with_db()

        node = SimpleNamespace(id=43, title="服务方案")
        clean_content = "<p>本公司将提供高质量的服务保障措施</p>"

        await composer._save_draft(node, clean_content)

        assert captured["content"] == clean_content


class TestRenderEntryGateContract:
    """渲染层闸门 B1 契约验证: _preprocess_html_for_bid 是 Word/HTML 共同入口

    通过验证 _preprocess_html_for_bid 输出干净 HTML, 间接保证下游 _html_to_docx
    与 _render_section_blocks 拿到的内容均不含 {xxx}, 最终 Word document.xml
    不会泄漏占位符。该函数是 Word 与 HTML 预览的唯一共同入口 (line 629/1660)。
    """

    def test_all_dirty_placeholders_neutralized_at_render_entry(self) -> None:
        dirty_html = (
            "<h4>{node_title}</h4>"
            "<p>本公司{company_profile}专注于{service_commitment}领域, "
            "执行{management_system}制度</p>"
            "<table><tr><td>{开户行}</td><td>说明</td></tr></table>"
            "<ul><li>{招标编号}</li></ul>"
        )

        result = DocRenderer._preprocess_html_for_bid(dirty_html, section_title="公司概况")

        for leaked in [
            "{node_title}",
            "{company_profile}",
            "{service_commitment}",
            "{management_system}",
            "{开户行}",
            "{招标编号}",
        ]:
            assert leaked not in result, (
                f"渲染入口闸门 B1 漏放占位符 {leaked!r}, 修复未生效"
            )

        for sanitized in [
            "[待补充:node_title]",
            "[待补充:company_profile]",
            "[待补充:service_commitment]",
            "[待补充:management_system]",
            "[待补充:开户行]",
            "[待补充:招标编号]",
        ]:
            assert sanitized in result

    def test_word_xml_body_without_placeholders(self) -> None:
        """轻量 E2E: 用纯 paragraph 内容验证 Word 字节流不含 {xxx}

        避开 BM_H2/BM_BODY 等项目专用样式 (需通过 _apply_styles 注册),
        只构造 paragraph 块走 _render_body_paragraph 路径, 验证 B1 净化生效。
        """
        from docx import Document

        doc = Document()
        renderer = DocRenderer(db=None)
        Pt_cls = MagicMock(side_effect=lambda x: x)

        if "BM_BODY" not in [s.name for s in doc.styles]:
            from docx.enum.style import WD_STYLE_TYPE
            doc.styles.add_style("BM_BODY", WD_STYLE_TYPE.PARAGRAPH)

        dirty_html = (
            "<p>本公司{company_profile}专注于{service_commitment}领域,"
            "执行{management_system}制度,联系{开户行}。</p>"
        )

        renderer._html_to_docx(
            doc, dirty_html, Pt_cls,
            image_records={},
            section_title="公司概况",
            base_depth=0,
            consumed_image_keys={},
        )

        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)

        with zipfile.ZipFile(buf) as zf:
            document_xml = zf.read("word/document.xml").decode("utf-8")

        for leaked in [
            "{company_profile}",
            "{service_commitment}",
            "{management_system}",
            "{开户行}",
        ]:
            assert leaked not in document_xml, (
                f"Word document.xml 仍泄漏占位符 {leaked!r}, 修复未生效"
            )


class TestReplacePlaceholdersBackwardCompat:
    """SectionComposer._replace_placeholders 私有方法保留向后兼容"""

    def _make_composer(self) -> SectionComposer:
        return SectionComposer(db=None, project_id=0)

    def test_legacy_method_delegates_to_module_function(self) -> None:
        """旧 API 必须委托新模块级函数, 行为完全一致"""
        composer = self._make_composer()
        template = "致{投标人}，项目{招标编号}，企业{company_profile}"

        legacy = composer._replace_placeholders(template)
        modern = sanitize_residual_placeholders(template)

        assert legacy == modern

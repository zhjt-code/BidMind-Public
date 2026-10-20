""": 写入 python-docx / 落库前剥离 XML 1.0 非法控制字符。

历史: 预览强制 Word 渲染；草稿含 \\x0c 等 C0 控制符时 lxml 抛
All strings must be XML compatible，全局异常处理器返回 500。
"""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from app.core.writer.doc_renderer import DocRenderer
from app.core.xml_compat import sanitize_xml_compatible_text

_XML_ERROR = "XML compatible"


def _docx_write(text: str) -> None:
    doc = Document()
    doc.add_paragraph().add_run(text)
    buffer = BytesIO()
    doc.save(buffer)


class TestSanitizeXmlCompatibleText:
    def test_strips_c0_controls_keeps_legal_whitespace_and_cjk(self) -> None:
        dirty = "工期\x00承诺\x0b共\x0c90天\x1f\t换行\n回车\r结束"
        result = sanitize_xml_compatible_text(dirty)

        assert "\x00" not in result
        assert "\x0b" not in result
        assert "\x0c" not in result
        assert "\x1f" not in result
        assert "\t" in result
        assert "\n" in result
        assert "\r" in result
        assert "工期承诺共90天" in result.replace("\t", "").replace("\n", "").replace("\r", "")

    def test_none_and_empty_and_idempotent(self) -> None:
        assert sanitize_xml_compatible_text(None) == ""
        assert sanitize_xml_compatible_text("") == ""
        clean = "正常正文"
        assert sanitize_xml_compatible_text(sanitize_xml_compatible_text(clean)) == clean


class TestPythonDocxRejectsIllegalXmlChars:
    def test_form_feed_raises_before_sanitize(self) -> None:
        with pytest.raises(ValueError, match=_XML_ERROR):
            _docx_write("第一章\x0c技术方案")

    def test_sanitized_text_can_save_docx(self) -> None:
        _docx_write(sanitize_xml_compatible_text("第一章\x0c技术方案\x1f完成"))


class TestWriterGatesStripIllegalXmlChars:
    def test_preprocess_html_strips_controls_from_historical_draft(self) -> None:
        dirty = "<p>质量保证\x0c体系</p><table><tr><td>A\x1f1</td></tr></table>"
        result = DocRenderer._preprocess_html_for_bid(dirty, section_title="质量保证")
        assert "\x0c" not in result
        assert "\x1f" not in result
        assert "质量保证体系" in result

    def test_body_paragraph_write_does_not_raise_on_form_feed(self) -> None:
        from docx.enum.style import WD_STYLE_TYPE

        doc = Document()
        if "BM_BODY" not in [s.name for s in doc.styles]:
            doc.styles.add_style("BM_BODY", WD_STYLE_TYPE.PARAGRAPH)
        DocRenderer._render_body_paragraph(doc, "实施方案\x0c如下")
        buffer = BytesIO()
        doc.save(buffer)
        assert buffer.tell() > 0

    @pytest.mark.asyncio
    async def test_save_draft_strips_controls_before_persist(self) -> None:
        import asyncio
        from types import SimpleNamespace
        from unittest.mock import AsyncMock

        from app.core.writer.composer.draft_persistence import save_draft

        class _ScalarRows:
            def scalars(self):
                return self

            def all(self):
                return []

        class _FakeSession:
            def __init__(self) -> None:
                self.added = None
                self.commit = AsyncMock()
                self.refresh = AsyncMock()

            async def execute(self, *_args, **_kwargs):
                return _ScalarRows()

            def add(self, obj):
                self.added = obj

        db = _FakeSession()
        composer = SimpleNamespace(
            db=db,
            _db_lock=asyncio.Lock(),
            _serialize_source_ids=lambda ids: "[]",
        )
        node = SimpleNamespace(id=9, title="质量保证")
        draft = await save_draft(composer, node, "<p>保证\x0c措施</p>")
        assert "\x0c" not in (draft.content or "")
        assert "保证措施" in (draft.content or "")

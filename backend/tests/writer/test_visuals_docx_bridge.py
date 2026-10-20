"""figure_bridge Word/预览插入测试（步骤4）。"""

from pathlib import Path
from types import SimpleNamespace

import pytest


@pytest.fixture
def diagram_png(tmp_path: Path) -> Path:
    from PIL import Image

    path = tmp_path / "diagram.png"
    Image.new("RGB", (400, 300), color=(240, 244, 248)).save(path)
    return path


def _record(diagram_id: int, png_path: str, title: str = "施工组织流程图"):
    return SimpleNamespace(
        id=diagram_id,
        png_path=png_path,
        title=title,
        diagram_type="flowchart",
        status="ready",
    )


class TestPreprocessRewrite:
    def test_doc_renderer_preprocess_rewrites_diagram_figure(self) -> None:
        from app.core.writer.doc_renderer import DocRenderer

        html = (
            "<p>正文段落</p>"
            '<figure data-diagram-id="7"><figcaption>施工组织流程图'
            "</figcaption></figure>"
        )
        result = DocRenderer._preprocess_html_for_bid(html, "施工组织方案")
        assert "[生成图示占位]" in result
        assert 'data-diagram-id="7"' in result
        assert "<figure" not in result


class TestInsertDiagramFigures:
    def test_inserts_png_into_placeholder(self, diagram_png: Path) -> None:
        from docx import Document
        from docx.shared import Pt

        from app.core.writer.visuals.figure_bridge import insert_diagram_figures

        doc = Document()
        doc.add_paragraph("正文段落")
        doc.add_paragraph("[生成图示占位]")
        raw_html = (
            '<figure data-diagram-id="7"><figcaption>施工组织流程图'
            "</figcaption></figure>"
        )
        insert_diagram_figures(
            doc, raw_html, Pt, {7: _record(7, str(diagram_png))},
        )
        assert not any("[生成图示占位]" in p.text for p in doc.paragraphs)
        assert any("施工组织流程图" in p.text for p in doc.paragraphs)
        assert len(doc.inline_shapes) == 1

    def test_missing_file_marks_pending(self, tmp_path: Path) -> None:
        from docx import Document
        from docx.shared import Pt

        from app.core.writer.visuals.figure_bridge import insert_diagram_figures

        doc = Document()
        doc.add_paragraph("[生成图示占位]")
        raw_html = (
            '<figure data-diagram-id="7"><figcaption>施工组织流程图'
            "</figcaption></figure>"
        )
        insert_diagram_figures(
            doc, raw_html, Pt,
            {7: _record(7, str(tmp_path / "not_exists.png"))},
        )
        texts = [p.text for p in doc.paragraphs]
        assert any("[图示待人工确认" in t for t in texts)
        assert len(doc.inline_shapes) == 0

    def test_tall_png_shrinks_below_nominal_width(self, tmp_path: Path) -> None:
        from docx import Document
        from docx.shared import Mm, Pt
        from PIL import Image

        from app.core.writer.visuals.figure_bridge import insert_diagram_figures
        from app.core.writer.visuals.page_fit import MAX_DIAGRAM_HEIGHT_MM

        png_path = tmp_path / "tall.png"
        Image.new("RGB", (400, 2000), color=(240, 244, 248)).save(png_path)
        doc = Document()
        doc.add_paragraph("[生成图示占位]")
        raw_html = (
            '<figure data-diagram-id="7"><figcaption>施工组织流程图'
            "</figcaption></figure>"
        )
        insert_diagram_figures(
            doc, raw_html, Pt, {7: _record(7, str(png_path))},
        )
        shape = doc.inline_shapes[0]
        assert int(shape.width) < int(Mm(130))
        expected_h = Mm(MAX_DIAGRAM_HEIGHT_MM)
        assert abs(int(shape.height) - int(expected_h)) < int(Mm(0.3))

    def test_unknown_diagram_id_marks_pending(self, diagram_png: Path) -> None:
        from docx import Document
        from docx.shared import Pt

        from app.core.writer.visuals.figure_bridge import insert_diagram_figures

        doc = Document()
        doc.add_paragraph("[生成图示占位]")
        raw_html = (
            '<figure data-diagram-id="99"><figcaption>未知图'
            "</figcaption></figure>"
        )
        insert_diagram_figures(
            doc, raw_html, Pt, {7: _record(7, str(diagram_png))},
        )
        texts = [p.text for p in doc.paragraphs]
        assert any("[图示待人工确认" in t for t in texts)


class TestResidualPlaceholders:
    def test_residual_marked_for_manual_review(self) -> None:
        from docx import Document

        from app.core.writer.visuals.figure_bridge import (
            mark_residual_diagram_placeholders,
        )

        doc = Document()
        doc.add_paragraph("[生成图示占位]")
        mark_residual_diagram_placeholders(doc)
        assert not any("[生成图示占位]" in p.text for p in doc.paragraphs)
        assert any("[图示待人工确认" in p.text for p in doc.paragraphs)


class TestPreviewInline:
    def test_inline_replaces_placeholder_with_data_uri(self, diagram_png: Path) -> None:
        from app.core.writer.visuals.figure_bridge import inline_diagram_figures

        rendered = "<p>正文</p><p>[生成图示占位]</p>"
        raw_html = (
            '<figure data-diagram-id="7"><figcaption>施工组织流程图'
            "</figcaption></figure>"
        )
        result = inline_diagram_figures(
            rendered, raw_html, {7: _record(7, str(diagram_png))},
        )
        assert "[生成图示占位]" not in result
        assert "data:image/png;base64," in result
        assert "施工组织流程图" in result
        assert "max-height:" in result
        assert "width:130.0mm" in result

    def test_inline_missing_record_marks_pending(self) -> None:
        from app.core.writer.visuals.figure_bridge import inline_diagram_figures

        rendered = "<p>[生成图示占位]</p>"
        raw_html = (
            '<figure data-diagram-id="7"><figcaption>施工组织流程图'
            "</figcaption></figure>"
        )
        result = inline_diagram_figures(rendered, raw_html, {})
        assert "[图示待人工确认" in result

    def test_no_diagram_returns_unchanged(self) -> None:
        from app.core.writer.visuals.figure_bridge import inline_diagram_figures

        rendered = "<p>正文</p>"
        assert inline_diagram_figures(rendered, "<p>正文</p>", {}) == rendered
